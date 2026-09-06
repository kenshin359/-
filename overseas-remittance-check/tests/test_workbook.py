#!/usr/bin/env python3
"""生成したExcelブックの数式が、実際に計算して正しく動くことのテスト。

数式は文字列として書き込まれるだけなので、「入っている」ことと「正しく動く」ことは別。
とくに照合列は、常に「一致」と返すだけの数式でも見た目は正常に見えてしまう。
そこで改ざんを検出できるかどうか（ネガティブテスト）まで確かめる。

純Pythonの評価器 formulas を使う。未導入ならスキップする。
評価に十数秒かかるため、他のテストより遅い。
"""

import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

try:
    import formulas  # noqa: F401
    HAVE_FORMULAS = True
except ImportError:
    HAVE_FORMULAS = False

SAMPLE = ROOT / "samples" / "remittance.pdf"


@unittest.skipUnless(HAVE_FORMULAS, "formulas が入っていない（pip install formulas）")
class TestWorkbookFormulas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import openpyxl
        from verify_workbook import evaluate

        cls.evaluate = staticmethod(evaluate)
        cls.tmp = tempfile.mkdtemp()
        # ASCII名にする。POSIXロケールでは日本語ファイル名が外部ツールに渡らない。
        cls.path = Path(cls.tmp) / "book.xlsx"

        cmd = [sys.executable, str(ROOT / "scripts" / "build_check_sheet.py"),
               str(ROOT / "data" / "OMT20260901100481.json"), "-o", str(cls.path)]
        if SAMPLE.exists():
            cmd += ["--pdf", str(SAMPLE)]
        subprocess.run(cmd, check=True, capture_output=True, cwd=ROOT)

        cls.wb = openpyxl.load_workbook(cls.path)
        cls.values = evaluate(cls.path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def get(self, sheet, coord):
        return self.values.get((sheet, coord))

    def test_full_calc_on_load_is_set(self):
        """これが無いと、開いた側が再計算するまで数式セルが空に見えることがある。"""
        self.assertTrue(self.wb.calculation.fullCalcOnLoad)

    def test_no_formula_errors(self):
        bad = [(s, c, v) for (s, c), v in self.values.items()
               if isinstance(v, str) and v.startswith("#")]
        self.assertEqual(bad, [], f"数式エラー: {bad}")

    def test_finding_counts_match_the_engine(self):
        import verify_remittance as vr
        import json
        tx = json.loads((ROOT / "data" / "OMT20260901100481.json").read_text(encoding="utf-8"))
        findings = vr.RemittanceVerifier(
            tx, vr.load_config("beneficiaries.json"), vr.load_config("bank_directory.json"),
            vr.load_config("risk_screening.json"), vr.load_config("products.json"),
        ).run()
        for coord, sev in (("B20", vr.BLOCK), ("B21", vr.WARN), ("B22", vr.INFO)):
            expected = sum(1 for f in findings if f.severity == sev)
            self.assertEqual(self.get("サマリー", coord), expected,
                             f"サマリー!{coord} の件数が検証エンジンと合っていない")

    def test_yen_estimates(self):
        amount = 238537.00
        self.assertAlmostEqual(self.get("サマリー", "C28"), amount * 145, places=2)
        self.assertAlmostEqual(self.get("サマリー", "C29"), amount * 150, places=2)
        self.assertAlmostEqual(self.get("サマリー", "C30"), amount * 155, places=2)
        # 基準(150)の行との差
        self.assertAlmostEqual(self.get("サマリー", "D28"), amount * (145 - 150), places=2)
        self.assertAlmostEqual(self.get("サマリー", "D30"), amount * (155 - 150), places=2)

    def test_all_reconciliation_rows_match(self):
        bad = [(c, v) for (s, c), v in self.values.items()
               if s == "照合結果" and c.startswith("D") and isinstance(v, str) and "❌" in v]
        self.assertEqual(bad, [], f"照合列に不一致がある: {bad}")

    def test_checklist_starts_at_zero_percent(self):
        self.assertAlmostEqual(self.get("実行前チェック", "D4"), 0.0, places=6)

    def test_tampering_is_detected(self):
        """常に「一致」と返すだけの数式になっていないことの確認。

        値を書き換えたら判定が反転しなければ、この照合表は何も検査していない。
        """
        import openpyxl
        tampered = Path(self.tmp) / "tampered.xlsx"
        shutil.copy(self.path, tampered)
        wb = openpyxl.load_workbook(tampered)
        ws = wb["照合結果"]
        ws["C9"] = "79969931316460"           # 口座番号の末尾1桁
        ws["C6"] = "HUGH TRADING CO LIMITED"  # 受取人名に語を挿入
        wb.save(tampered)

        values = self.evaluate(tampered)
        self.assertIn("❌", str(values.get(("照合結果", "D9"))), "口座番号の差し替えを検出できていない")
        self.assertIn("❌", str(values.get(("照合結果", "D6"))), "受取人名の変更を検出できていない")
        # 触っていない行は一致のまま
        for coord in ("D7", "D10", "D11"):
            self.assertIn("✅", str(values.get(("照合結果", coord))),
                          f"{coord} が巻き添えで不一致になっている")
