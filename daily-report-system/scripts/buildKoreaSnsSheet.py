#!/usr/bin/env python3
# ============================================================
#  韓国SNS（明洞）管理表（Excel）の組み立て
# ------------------------------------------------------------
#  韓国チームが毎日書き込む「撮影 → 編集 → 投稿 → 広告 → 予約」の管理表。
#  kintoneアプリ「韓国SNS運用管理（明洞）」と同じ項目・同じ判定なので、
#  Excelで運用を始めて、あとからkintoneへ移しても中身が変わりません。
#
#  シート:
#    サマリー            … アカウント別の実績と、いま止まっている件数
#    ① 撮影スケジュール   … 撮影予定（事前報告）。2日前ルールを自動判定
#    ② 投稿スケジュール   … 1投稿＝1行。投稿予定日の遅れを自動判定
#    ③ 案件進捗          … 事前→撮影→編集→投稿→広告→流入→予約を1行で追う
#                          （①②から自動で集まる。手で入れるのは黄色のセルだけ）
#    選択肢              … プルダウンの元（アカウント・媒体・形態など）
#
#  実行:
#    python3 scripts/buildKoreaSnsSheet.py                    … 空の管理表（記入例つき）
#    python3 scripts/buildKoreaSnsSheet.py --month=2026-09
#    python3 scripts/buildKoreaSnsSheet.py --json=out/korea-sns.json
#        ↑ `npm run korea:sns -- --json=out/korea-sns.json` の結果を流し込む
#
#  ★数式で計算させています（手で計算した数字は書きません）。
#    行を足したり日付を直したりすれば、判定も合計も自動で追従します。
# ============================================================
import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG = os.path.join(ROOT, 'config', 'korea-sns.json')

F = 'Yu Gothic'
NUM = '#,##0'
YEN = '¥#,##0;(¥#,##0);-'
DATE_FMT = 'yyyy/mm/dd'

HDR_FILL = PatternFill('solid', fgColor='188038')       # 見出し（緑）
INPUT_FILL = PatternFill('solid', fgColor='FFF7CC')     # 手で入れる欄（黄）
AUTO_FILL = PatternFill('solid', fgColor='EFEFEF')      # 自動計算の欄（灰）
SUM_FILL = PatternFill('solid', fgColor='E2EFDA')
NG_FILL = PatternFill('solid', fgColor='FCE4E4')        # 🔴 の行

HDR_FONT = Font(name=F, bold=True, color='FFFFFF', size=10)
BOLD = Font(name=F, bold=True)
NORM = Font(name=F, size=10)
SAMPLE = Font(name=F, size=10, italic=True, color='808080')
NOTE = Font(name=F, size=9, color='808080')
TITLE = Font(name=F, bold=True, size=14)
THIN = Border(*[Side(style='thin', color='BFBFBF')] * 4)

HEAD_ROW = 4       # 見出しの行
FIRST_ROW = 5      # 記入例の行（＝データの1行目）
LAST_ROW = 204     # 用意しておく行数（200行）

S_SHOOT = '① 撮影スケジュール'
S_POST = '② 投稿スケジュール'
S_CASE = '③ 案件進捗'
S_OPT = '選択肢'


def arg(name, default=None):
    for a in sys.argv[1:]:
        if a.startswith(f'--{name}='):
            return a[len(name) + 3:]
    return default


def load_config():
    with open(CONFIG, encoding='utf-8') as f:
        return json.load(f)


def head(ws, row, cols, widths):
    """見出し行を作る。cols は (列名, 入力/自動) のリスト。"""
    for i, (label, kind) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.border = THIN
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def title_block(ws, title, note):
    ws['A1'] = title
    ws['A1'].font = TITLE
    ws['A2'] = note
    ws['A2'].font = NOTE
    ws['A3'] = '■ 黄色のセル＝手で入力 / 灰色のセル＝自動計算（触らない）'
    ws['A3'].font = NOTE


def style_rows(ws, cols, first=FIRST_ROW, last=LAST_ROW):
    """行の見た目（入力欄は黄色・自動欄は灰色）をまとめて付ける。"""
    for r in range(first, last + 1):
        for i, (_, kind) in enumerate(cols, start=1):
            c = ws.cell(row=r, column=i)
            c.font = SAMPLE if r == first else NORM
            c.border = THIN
            c.fill = INPUT_FILL if kind == 'in' else AUTO_FILL
            c.alignment = Alignment(vertical='center')


def fmt_col(ws, col, number_format, first=FIRST_ROW, last=LAST_ROW):
    for r in range(first, last + 1):
        ws.cell(row=r, column=col).number_format = number_format


def dv(ws, sheet_opt_range, col, first=FIRST_ROW, last=LAST_ROW):
    """プルダウンを付ける（選択肢シートを参照）。"""
    v = DataValidation(type='list', formula1=sheet_opt_range, allow_blank=True, showDropDown=False)
    ws.add_data_validation(v)
    v.add(f'{get_column_letter(col)}{first}:{get_column_letter(col)}{last}')


def ng_rule(ws, judge_col, cols_count, first=FIRST_ROW, last=LAST_ROW):
    """判定に 🔴 が出た行を赤くする。"""
    rng = f'A{first}:{get_column_letter(cols_count)}{last}'
    col = get_column_letter(judge_col)
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=[f'ISNUMBER(SEARCH("🔴",${col}{first}))'], fill=NG_FILL, stopIfTrue=False),
    )


# ────────────────────────────────────────────────
#  選択肢シート
# ────────────────────────────────────────────────
def build_options(wb, cfg):
    ws = wb.create_sheet(S_OPT)
    title_block(ws, '選択肢マスタ', 'プルダウンの元です。増やしたいときはここに足してください（列の順番は変えないこと）。')
    lists = [
        ('アカウント', [a['id'] for a in cfg['accounts']]),
        ('媒体', cfg['media']),
        ('投稿形態', cfg['formats']),
        ('投稿の種類', cfg['post_types']),
        ('撮影の目的', cfg['goals']),
        ('進捗', ['① 事前共有済み', '② 撮影済み', '③ 編集中', '④ 投稿済み', '⑤ 広告運用中', '⑥ 成果集計済み', '中止']),
        ('広告利用', ['あり', 'なし', '未定']),
    ]
    for i, (name, values) in enumerate(lists, start=1):
        c = ws.cell(row=HEAD_ROW, column=i, value=name)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, THIN
        ws.column_dimensions[get_column_letter(i)].width = 20
        for j, v in enumerate(values, start=FIRST_ROW):
            cell = ws.cell(row=j, column=i, value=v)
            cell.font, cell.border = NORM, THIN

    # アカウントのプロフィール（どの店のアカウントか迷わないように）
    r = FIRST_ROW + 12
    ws.cell(row=r, column=1, value='運用アカウント').font = BOLD
    for k, a in enumerate(cfg['accounts'], start=1):
        if not a.get('url'):
            continue
        ws.cell(row=r + k, column=1, value=a['id']).font = NORM
        ws.cell(row=r + k, column=2, value=a['name']).font = NORM
        link = ws.cell(row=r + k, column=3, value=a['url'])
        link.font = Font(name=F, size=10, color='0563C1', underline='single')
        link.hyperlink = a['url']
    return ws


def opt_range(col, count):
    # ★先頭に「=」を付けないこと（付けるとExcelがプルダウンを読めない）
    return f"'{S_OPT}'!${col}${FIRST_ROW}:${col}${FIRST_ROW + count - 1}"


# ────────────────────────────────────────────────
#  ① 撮影スケジュール
# ────────────────────────────────────────────────
SHOOT_COLS = [
    ('案件ID', 'in'), ('撮影日', 'in'), ('曜日', 'auto'), ('時間', 'in'), ('アカウント', 'in'),
    ('案件名', 'in'), ('撮影場所', 'in'), ('撮影内容', 'in'), ('撮影の目的', 'in'),
    ('出演者・インフルエンサー', 'in'), ('フォロワー数', 'in'), ('撮影予定本数', 'in'),
    ('投稿予定媒体', 'in'), ('投稿予定日', 'in'), ('広告利用', 'in'),
    ('事前共有した日', 'in'), ('何日前に共有', 'auto'), ('★2日前ルール', 'auto'),
    ('案件担当', 'in'), ('編集担当', 'in'),
]
SHOOT_W = [10, 12, 6, 14, 11, 26, 16, 26, 14, 20, 11, 11, 16, 12, 9, 13, 10, 17, 11, 11]


def build_shoot(wb, cfg, rows):
    ws = wb.create_sheet(S_SHOOT)
    title_block(
        ws, '① 撮影スケジュール（事前報告）',
        f'撮影が決まったらすぐ1行。事前共有は【撮影日の{cfg["rules"]["advance_notice_days"]}日前まで】。'
        'R列が🔴になった撮影はルール違反として週次レポートに出ます。',
    )
    head(ws, HEAD_ROW, SHOOT_COLS, SHOOT_W)
    style_rows(ws, SHOOT_COLS)

    n = cfg['rules']['advance_notice_days']
    for r in range(FIRST_ROW, LAST_ROW + 1):
        ws.cell(row=r, column=3, value=(
            f'=IF($B{r}="","",CHOOSE(WEEKDAY($B{r}),"日","月","火","水","木","金","土"))'))
        ws.cell(row=r, column=17, value=f'=IF(OR($B{r}="",$P{r}=""),"",$B{r}-$P{r})')
        ws.cell(row=r, column=18, value=(
            f'=IF($B{r}="","",'
            f'IF($P{r}="","🔴 事前共有なし",'
            f'IF($Q{r}>={n},"🟢 OK","🔴 "&$Q{r}&"日前")))'))
        ws.cell(row=r, column=2).number_format = DATE_FMT
        ws.cell(row=r, column=14).number_format = DATE_FMT
        ws.cell(row=r, column=16).number_format = DATE_FMT
        ws.cell(row=r, column=11).number_format = NUM
    ng_rule(ws, 18, len(SHOOT_COLS))

    dv(ws, opt_range('A', len(cfg['accounts'])), 5)
    dv(ws, opt_range('B', len(cfg['media'])), 13)
    dv(ws, opt_range('E', len(cfg['goals'])), 9)
    dv(ws, opt_range('G', 3), 15)

    sample = ['T-001', date(2026, 9, 20), None, '14:00〜17:00', 'Tiffany',
              '9/20 明洞店 インフルエンサー撮影', '明洞本店', '来店〜施術の流れ', '予約獲得',
              '@myeongdong_kim', 52000, 3, 'Instagram', date(2026, 9, 25), 'あり',
              date(2026, 9, 17), None, None, 'ミンジ', 'ソンチャン']
    fill_rows(ws, [sample] if not rows else rows, skip_cols={3, 17, 18}, sample=not rows)
    if not rows:
        ws.cell(row=FIRST_ROW, column=len(SHOOT_COLS) + 1, value='← 記入例。使うときは消してください').font = NOTE
    return ws


# ────────────────────────────────────────────────
#  ② 投稿スケジュール
# ────────────────────────────────────────────────
POST_COLS = [
    ('案件ID', 'in'), ('アカウント', 'in'), ('案件名', 'in'), ('投稿予定日', 'in'), ('実投稿日', 'in'),
    ('媒体', 'in'), ('投稿形態', 'in'), ('投稿の種類', 'in'), ('投稿URL', 'in'),
    ('広告利用', 'in'), ('広告費', 'in'), ('再生・表示数', 'in'), ('いいね', 'in'), ('保存', 'in'),
    ('プロフィール遷移', 'in'), ('リンククリック', 'in'), ('★状態', 'auto'),
    ('ひとこと（伸びた/落ちた理由）', 'in'),
]
POST_W = [10, 11, 24, 12, 12, 12, 12, 20, 34, 9, 12, 12, 9, 9, 14, 13, 15, 30]
POST_LAST = FIRST_ROW + 399  # 投稿は案件より多いので400行


def build_post(wb, cfg, rows):
    ws = wb.create_sheet(S_POST)
    title_block(
        ws, '② 投稿スケジュール（1投稿＝1行）',
        '投稿する予定を先に書き、投稿したら「実投稿日」とURLを入れます。'
        '案件IDは①と同じものを入れてください（③の集計がこのIDでつながります）。',
    )
    head(ws, HEAD_ROW, POST_COLS, POST_W)
    style_rows(ws, POST_COLS, last=POST_LAST)

    late = cfg['rules']['post_overdue_days']
    for r in range(FIRST_ROW, POST_LAST + 1):
        ws.cell(row=r, column=17, value=(
            f'=IF($A{r}="","",'
            f'IF($E{r}<>"",IF($I{r}="","🟡 URL未記入","🟢 投稿済み"),'
            f'IF($D{r}="","🟡 予定日未定",'
            f'IF(TODAY()-$D{r}>{late},"🔴 "&TEXT(TODAY()-$D{r},"0")&"日遅れ","🟡 投稿まち"))))'))
        for col in (4, 5):
            ws.cell(row=r, column=col).number_format = DATE_FMT
        ws.cell(row=r, column=11).number_format = YEN
        for col in (12, 13, 14, 15, 16):
            ws.cell(row=r, column=col).number_format = NUM
    ng_rule(ws, 17, len(POST_COLS), last=POST_LAST)

    dv(ws, opt_range('A', len(cfg['accounts'])), 2, last=POST_LAST)
    dv(ws, opt_range('B', len(cfg['media'])), 6, last=POST_LAST)
    dv(ws, opt_range('C', len(cfg['formats'])), 7, last=POST_LAST)
    dv(ws, opt_range('D', len(cfg['post_types'])), 8, last=POST_LAST)
    dv(ws, opt_range('G', 2), 10, last=POST_LAST)

    sample = ['T-001', 'Tiffany', '9/20 明洞店 インフルエンサー撮影', date(2026, 9, 25),
              date(2026, 9, 25), 'Instagram', 'リール', 'タイアップ（PR表記あり）',
              'https://www.instagram.com/reel/xxxxxxxx/', 'あり', 120000, 98000, 3200, 410, 850, 320,
              None, '冒頭2秒で施術シーンを出したら最後まで見られた']
    fill_rows(ws, [sample] if not rows else rows, skip_cols={17}, last=POST_LAST, sample=not rows)
    if not rows:
        ws.cell(row=FIRST_ROW, column=len(POST_COLS) + 1, value='← 記入例。使うときは消してください').font = NOTE
    return ws


# ────────────────────────────────────────────────
#  ③ 案件進捗（インフルエンサー案件 → 投稿管理まで）
# ────────────────────────────────────────────────
CASE_COLS = [
    ('案件ID', 'auto'), ('アカウント', 'auto'), ('案件名', 'auto'), ('出演者', 'auto'),
    ('撮影日', 'auto'), ('事前共有日', 'auto'), ('★2日前ルール', 'auto'),
    ('撮影本数（実績）', 'in'), ('編集予定本数', 'in'), ('編集完了予定日', 'in'), ('編集完了日', 'in'),
    ('投稿本数', 'auto'), ('最終投稿日', 'auto'), ('再生・表示数', 'auto'), ('広告費', 'auto'),
    ('リンククリック', 'auto'),
    ('SNS経由の流入', 'in'), ('予約件数', 'in'), ('来店件数', 'in'), ('売上（税込）', 'in'),
    ('予約単価', 'auto'), ('進捗', 'in'), ('★次にやること', 'auto'),
]
CASE_W = [10, 11, 26, 18, 12, 12, 17, 13, 12, 14, 12, 10, 12, 13, 12, 13, 13, 11, 11, 14, 12, 15, 30]


def build_case(wb, cfg):
    ws = wb.create_sheet(S_CASE)
    title_block(
        ws, '③ 案件進捗（撮影 → 編集 → 投稿 → 広告 → 流入 → 予約）',
        '①②に書けば、白紙のところ以外は自動で集まります。手で入れるのは黄色（撮影実績・編集・流入・予約・売上）だけ。'
        'W列「次にやること」が、その案件でいま止まっていることです。',
    )
    head(ws, HEAD_ROW, CASE_COLS, CASE_W)
    style_rows(ws, CASE_COLS)

    res_days = cfg['rules']['result_days_after_post']
    edit_days = cfg['rules']['edit_overdue_days']
    post_days = cfg['rules']['post_overdue_days']
    target_cpa = cfg['targets']['cpa_yen']
    P = f"'{S_POST}'"
    S = f"'{S_SHOOT}'"

    def prng(col):
        return f'{P}!${col}{FIRST_ROW}:${col}{POST_LAST}'

    for r in range(FIRST_ROW, LAST_ROW + 1):
        # ①（撮影スケジュール）と同じ行を見るだけ。IDの突き合わせは不要。
        ws.cell(row=r, column=1, value=f'=IF({S}!$A{r}="","",{S}!$A{r})')
        ws.cell(row=r, column=2, value=f'=IF({S}!$E{r}="","",{S}!$E{r})')
        ws.cell(row=r, column=3, value=f'=IF({S}!$F{r}="","",{S}!$F{r})')
        ws.cell(row=r, column=4, value=f'=IF({S}!$J{r}="","",{S}!$J{r})')
        ws.cell(row=r, column=5, value=f'=IF({S}!$B{r}="","",{S}!$B{r})')
        ws.cell(row=r, column=6, value=f'=IF({S}!$P{r}="","",{S}!$P{r})')
        ws.cell(row=r, column=7, value=f'=IF({S}!$R{r}="","",{S}!$R{r})')

        # ②（投稿スケジュール）から案件IDで集める
        # 投稿本数＝案件IDが一致し、実投稿日が入っている行の数
        #   ★COUNTIFS の "<>"（空欄でない）は表計算ソフトによって解釈が割れるため使わない
        ws.cell(row=r, column=12, value=(
            f'=IF($A{r}="","",SUMPRODUCT(({prng("A")}=$A{r})*({prng("E")}<>"")))'))
        # 最終投稿日。MAXIFS は新しいExcelにしか無いので SUMPRODUCT(MAX(…)) で代用する
        ws.cell(row=r, column=13, value=(
            f'=IF($L{r}=0,"",SUMPRODUCT(MAX(({prng("A")}=$A{r})*{prng("E")})))'))
        ws.cell(row=r, column=14, value=f'=IF($A{r}="","",SUMIFS({prng("L")},{prng("A")},$A{r}))')
        ws.cell(row=r, column=15, value=f'=IF($A{r}="","",SUMIFS({prng("K")},{prng("A")},$A{r}))')
        ws.cell(row=r, column=16, value=f'=IF($A{r}="","",SUMIFS({prng("P")},{prng("A")},$A{r}))')

        # 予約単価（0除算しない）
        ws.cell(row=r, column=21, value=f'=IF(OR($R{r}="",$R{r}=0),"",$O{r}/$R{r})')

        # いま止まっていること（上から順に優先度が高い）
        ws.cell(row=r, column=23, value=(
            f'=IF($A{r}="","",'
            f'IF($V{r}="中止","—",'
            f'IF(ISNUMBER(SEARCH("🔴",$G{r})),"🔴 事前共有ルール違反",'
            f'IF(AND($E{r}<>"",$E{r}<TODAY(),$H{r}=""),"🔴 撮影後の報告まち",'
            # ★Excelの AND は中身を全部計算するので、空欄との引き算を IFERROR で守る
            f'IF(AND($K{r}="",$J{r}<>"",IFERROR(TODAY()-$J{r}>{edit_days},FALSE)),"🔴 編集が遅れています（"&$A{r}&"）",'
            f'IF(AND($L{r}=0,{S}!$N{r}<>"",IFERROR(TODAY()-{S}!$N{r}>{post_days},FALSE)),"🔴 投稿まち",'
            f'IF(AND($L{r}>0,OR($Q{r}="",$R{r}=""),IFERROR(TODAY()-$M{r}>{res_days},FALSE)),"🔴 流入・予約が未入力",'
            f'IF($R{r}<>"","🟢 成果まで記入ずみ","🟡 進行中"))))))))'))

        for col in (5, 6, 10, 11, 13):
            ws.cell(row=r, column=col).number_format = DATE_FMT
        for col in (8, 9, 12, 14, 16, 17, 18, 19):
            ws.cell(row=r, column=col).number_format = NUM
        for col in (15, 20, 21):
            ws.cell(row=r, column=col).number_format = YEN

    ng_rule(ws, 23, len(CASE_COLS))
    # 予約単価が目標を超えた案件を赤字にする
    ws.conditional_formatting.add(
        f'U{FIRST_ROW}:U{LAST_ROW}',
        FormulaRule(formula=[f'AND($U{FIRST_ROW}<>"",$U{FIRST_ROW}>{target_cpa})'],
                    font=Font(name=F, color='C00000', bold=True), stopIfTrue=False),
    )
    dv(ws, opt_range('F', 7), 22)
    return ws


# ────────────────────────────────────────────────
#  サマリー
# ────────────────────────────────────────────────
def build_summary(wb, cfg, month):
    ws = wb.create_sheet('サマリー', 0)
    title_block(
        ws, f'韓国SNS（明洞）管理表　{month}',
        '①②③に書き込むと、ここが自動で更新されます。数字はすべて数式です（手計算は入っていません）。',
    )
    C = f"'{S_CASE}'"

    def crng(col):
        return f'{C}!${col}${FIRST_ROW}:${col}${LAST_ROW}'

    cols = ['アカウント', '案件数', '撮影本数', '投稿本数', '再生・表示数', '広告費',
            '流入', '予約', '来店', '売上', '予約単価', '判定']
    widths = [14, 9, 10, 9, 13, 13, 10, 9, 9, 14, 12, 22]
    head(ws, HEAD_ROW, [(c, 'auto') for c in cols], widths)

    target = cfg['targets']['cpa_yen']
    accounts = [a['id'] for a in cfg['accounts']]
    r = FIRST_ROW
    for acc in accounts:
        ws.cell(row=r, column=1, value=acc).font = BOLD
        ws.cell(row=r, column=2, value=f'=COUNTIFS({crng("B")},$A{r})')
        ws.cell(row=r, column=3, value=f'=SUMIFS({crng("H")},{crng("B")},$A{r})')
        ws.cell(row=r, column=4, value=f'=SUMIFS({crng("L")},{crng("B")},$A{r})')
        ws.cell(row=r, column=5, value=f'=SUMIFS({crng("N")},{crng("B")},$A{r})')
        ws.cell(row=r, column=6, value=f'=SUMIFS({crng("O")},{crng("B")},$A{r})')
        ws.cell(row=r, column=7, value=f'=SUMIFS({crng("Q")},{crng("B")},$A{r})')
        ws.cell(row=r, column=8, value=f'=SUMIFS({crng("R")},{crng("B")},$A{r})')
        ws.cell(row=r, column=9, value=f'=SUMIFS({crng("S")},{crng("B")},$A{r})')
        ws.cell(row=r, column=10, value=f'=SUMIFS({crng("T")},{crng("B")},$A{r})')
        ws.cell(row=r, column=11, value=f'=IF($H{r}=0,"",$F{r}/$H{r})')
        ws.cell(row=r, column=12, value=(
            f'=IF($K{r}="","予約が未入力",'
            f'IF($K{r}<={target},"🟢 目標内","🔴 目標¥{target:,} 超過"))'))
        r += 1

    total = r
    ws.cell(row=total, column=1, value='合計').font = BOLD
    for col in range(2, 11):
        L = get_column_letter(col)
        ws.cell(row=total, column=col, value=f'=SUM({L}{FIRST_ROW}:{L}{total - 1})')
    ws.cell(row=total, column=11, value=f'=IF($H{total}=0,"",$F{total}/$H{total})')
    ws.cell(row=total, column=12, value=(
        f'=IF($K{total}="","予約が未入力",IF($K{total}<={target},"🟢 目標内","🔴 目標¥{target:,} 超過"))'))

    for rr in range(FIRST_ROW, total + 1):
        for col in range(1, len(cols) + 1):
            c = ws.cell(row=rr, column=col)
            c.font = BOLD if rr == total else NORM
            c.border = THIN
            if rr == total:
                c.fill = SUM_FILL
        for col in (2, 3, 4, 5, 7, 8, 9):
            ws.cell(row=rr, column=col).number_format = NUM
        for col in (6, 10, 11):
            ws.cell(row=rr, column=col).number_format = YEN

    # いま止まっているもの（🔴 の件数）
    r = total + 2
    ws.cell(row=r, column=1, value='いま止まっているもの').font = BOLD
    items = [
        ('🔴 要対応（案件）', f'=COUNTIF({crng("W")},"🔴*")'),
        ('🟡 進行中（案件）', f'=COUNTIF({crng("W")},"🟡*")'),
        ('🟢 成果まで記入ずみ', f'=COUNTIF({crng("W")},"🟢*")'),
        ('🔴 投稿の遅れ（投稿）', f'=COUNTIF(\'{S_POST}\'!$Q${FIRST_ROW}:$Q${POST_LAST},"🔴*")'),
        ('🔴 事前共有ルール違反（撮影）', f'=COUNTIF(\'{S_SHOOT}\'!$R${FIRST_ROW}:$R${LAST_ROW},"🔴*")'),
    ]
    for k, (label, formula) in enumerate(items, start=1):
        ws.cell(row=r + k, column=1, value=label).font = NORM
        c = ws.cell(row=r + k, column=2, value=formula)
        c.font, c.border, c.number_format = BOLD, THIN, NUM
        c.fill = INPUT_FILL if label.startswith('🔴') else AUTO_FILL

    r = r + len(items) + 2
    ws.cell(row=r, column=1, value='目標（config/korea-sns.json）').font = BOLD
    targets = [
        ('月の撮影件数', cfg['targets']['shoots_per_month']),
        ('月の投稿本数', cfg['targets']['posts_per_month']),
        ('月の予約件数', cfg['targets']['reservations_per_month']),
        ('予約1件あたりの広告費（上限）', cfg['targets']['cpa_yen']),
    ]
    for k, (label, v) in enumerate(targets, start=1):
        ws.cell(row=r + k, column=1, value=label).font = NORM
        c = ws.cell(row=r + k, column=2, value=v)
        c.font, c.border = NORM, THIN
        c.number_format = YEN if '広告費' in label else NUM

    r = r + len(targets) + 2
    for k, line in enumerate([
        '使い方: ①に撮影予定 → ②に投稿予定 → 投稿したら②に実績 → ③の黄色（流入・予約・売上）を埋める',
        '事前共有は撮影日の2日前まで。①のR列が🔴の撮影はルール違反として日本側に共有されます。',
        'kintoneアプリ「韓国SNS運用管理（明洞）」と同じ項目です。あとから移しても書き直しは不要です。',
    ]):
        ws.cell(row=r + k, column=1, value=line).font = NOTE
    ws.freeze_panes = 'A5'
    return ws


def fill_rows(ws, rows, skip_cols=frozenset(), first=FIRST_ROW, last=LAST_ROW, sample=False):
    """データ行を書く（数式の列は飛ばす）。sample=True のときは記入例として灰色斜体にする。"""
    for i, row in enumerate(rows):
        r = first + i
        if r > last:
            break
        for j, v in enumerate(row, start=1):
            if j in skip_cols or v is None or v == '':
                continue
            c = ws.cell(row=r, column=j, value=v)
            c.font = SAMPLE if sample else NORM


def to_date(v):
    """'2026-09-20' のような文字列を日付にする（日付でないものはそのまま返す）。"""
    if isinstance(v, str) and len(v) == 10 and v[4] == '-' and v[7] == '-':
        try:
            y, m, d = (int(x) for x in v.split('-'))
            return date(y, m, d)
        except ValueError:
            return v
    return v


def rows_from_json(path):
    """`npm run korea:sns -- --json=…` の集計から ①② の行を作る。"""
    with open(path, encoding='utf-8') as f:
        summary = json.load(f)
    shoots, posts = [], []
    for i, it in enumerate(summary.get('items', []), start=1):
        case_id = it.get('id') or f'A-{i:03d}'
        shoots.append([
            case_id, it.get('shootDate'), None, it.get('shootTime'), it.get('account'),
            it.get('title'), it.get('place'), it.get('content'), it.get('goal'),
            it.get('cast'), it.get('castFollower'), it.get('planCount'),
            '・'.join(it.get('planMedia') or []), it.get('planPostDate'), it.get('planAd'),
            it.get('sharedAt'), None, None, it.get('owner'), it.get('editor'),
        ])
        for p in it.get('posts', []):
            posts.append([
                case_id, it.get('account'), it.get('title'), it.get('planPostDate'), p.get('postDate'),
                p.get('media'), p.get('format'), p.get('type'), p.get('url'),
                'あり' if p.get('ad') else 'なし', p.get('adCost'), p.get('views'), p.get('likes'),
                p.get('saves'), p.get('profile'), p.get('clicks'), None, p.get('note'),
            ])
    return ([[to_date(v) for v in row] for row in shoots],
            [[to_date(v) for v in row] for row in posts])


def main():
    cfg = load_config()
    month = arg('month', date.today().strftime('%Y-%m'))
    src = arg('json')
    shoots, posts = rows_from_json(src) if src else ([], [])

    wb = Workbook()
    wb.remove(wb.active)
    # 開いた瞬間に全部の数式を計算させる（openpyxl は計算結果を持たないため）
    wb.calculation.fullCalcOnLoad = True
    build_options(wb, cfg)
    build_shoot(wb, cfg, shoots)
    build_post(wb, cfg, posts)
    build_case(wb, cfg)
    build_summary(wb, cfg, month)
    # 並び順: サマリー → ① → ② → ③ → 選択肢
    order = ['サマリー', S_SHOOT, S_POST, S_CASE, S_OPT]
    wb._sheets = [wb[name] for name in order]

    out = arg('out') or os.path.join(ROOT, 'out', f'韓国SNS管理表_{month}.xlsx')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f'作成しました: {out}')
    if src:
        print(f'  （{src} から 撮影{len(shoots)}件 / 投稿{len(posts)}本 を流し込みました）')
    else:
        print('  （空の管理表です。1行目は記入例なので、使うときは消してください）')


if __name__ == '__main__':
    main()
