#!/usr/bin/env python3
# ============================================================
#  会員手続き書類（入会同意書・休会届・退会届）の生成
# ------------------------------------------------------------
#  パーソナルジムの窓口で使う3枚の用紙と、受付の記録台帳・締切ルールを
#  1つのExcelにまとめます。A4たてで印刷してそのまま記入できます。
#
#  入力: config/member-forms.json（文面・プラン・締切はすべてここ）
#  出力: out/会員手続き書類_<店舗>.xlsx
#    ① 入会同意書   ② 休会届   ③ 退会届
#    ④ 受付台帳（手続きの記録・処理もれの防止）
#    ⑤ 手続きルール（締切と社内フロー）
#    ⑥ 法務チェック（専門家に見てもらう前の確認事項）
#
#  実行:
#    npm run forms:member
#    python3 scripts/buildMemberForms.py --store=梅田店
#
#  ★この書式はひな形です。◯や（要確認）の箇所を埋め、最終文面は
#    弁護士等の確認を受けてから配布してください（⑥シートに一覧）。
# ============================================================
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlsx_common import setup_print  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

F = 'Yu Gothic'
TITLE = Font(name=F, bold=True, size=16)
SUB = Font(name=F, size=9.5, color='595959')
H2 = Font(name=F, bold=True, size=11)
HEAD = Font(name=F, bold=True, color='FFFFFF', size=10.5)
BODY = Font(name=F, size=10.5)
SMALL = Font(name=F, size=9.8)
TINY = Font(name=F, size=8.5, color='595959')
BOLD = Font(name=F, bold=True, size=10.5)
RED = Font(name=F, bold=True, size=10.5, color='C00000')

NAVY = PatternFill('solid', fgColor='1F3864')
LABEL = PatternFill('solid', fgColor='F2F2F2')
WARN = PatternFill('solid', fgColor='FFF2CC')
SIGN = PatternFill('solid', fgColor='FDF2F2')

thin = Side(style='thin', color='808080')
hair = Side(style='hair', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

COLS = 8            # A〜H で1枚を組む
WIDTHS = [13, 11, 11, 11, 11, 11, 11, 12]
CHECK = '☐ '


def arg(name, fallback=None):
    hit = next((a for a in sys.argv[1:] if a.startswith(f'--{name}=')), None)
    return hit.split('=', 1)[1] if hit else fallback


def setup(ws, landscape=False):
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    setup_print(ws, landscape=landscape, centered=True)


def span(ws, row, c1, c2, value=None, font=BODY, fill=None, border=None,
         wrap=False, align='left', valign='center', height=None):
    """c1〜c2列を結合して1つの記入欄・見出しにする。"""
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    cell = ws.cell(row=row, column=c1, value=value)
    cell.font = font
    cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    for c in range(c1, c2 + 1):
        cc = ws.cell(row=row, column=c)
        if fill:
            cc.fill = fill
        if border:
            cc.border = border
    if height:
        ws.row_dimensions[row].height = height
    return cell


def band(ws, row, text):
    """紺色の見出し帯。"""
    span(ws, row, 1, COLS, text, HEAD, NAVY, BOX, height=20)
    return row + 1


def field(ws, row, label, value='', label_to=2, height=26, font=BODY):
    """左に項目名（灰色）、右に記入欄（枠線）。"""
    span(ws, row, 1, label_to, label, BOLD, LABEL, BOX, align='left')
    span(ws, row, label_to + 1, COLS, value, font, None, BOX, align='left', wrap=True)
    ws.row_dimensions[row].height = height
    return row + 1


def pair(ws, row, l1, l2, height=26):
    """1行に2項目（左半分・右半分）。"""
    span(ws, row, 1, 2, l1, BOLD, LABEL, BOX)
    span(ws, row, 3, 4, '', BODY, None, BOX)
    span(ws, row, 5, 5, l2, BOLD, LABEL, BOX)
    span(ws, row, 6, COLS, '', BODY, None, BOX)
    ws.row_dimensions[row].height = height
    return row + 1


def para(ws, row, text, font=SMALL, per_line=58, indent=1):
    """折り返しの説明文。行数から高さを決める。"""
    lines = max(1, -(-len(text) // per_line))
    span(ws, row, indent, COLS, text, font, None, None, wrap=True, align='left',
         valign='top', height=13.5 * lines + 4)
    return row + 1


def checkboxes(ws, row, items, per_row=2):
    """□つきの選択肢を並べる。"""
    step = COLS // per_row
    i = 0
    while i < len(items):
        for j in range(per_row):
            if i + j >= len(items):
                break
            c1 = 1 + j * step
            c2 = c1 + step - 1
            span(ws, row, c1, c2, CHECK + items[i + j], SMALL, None, None, wrap=True)
        ws.row_dimensions[row].height = 20
        row += 1
        i += per_row
    return row


# ------------------------------------------------------------
#  ① 入会同意書
# ------------------------------------------------------------
def sheet_join(wb, cfg):
    ws = wb.create_sheet('①入会同意書')
    setup(ws)
    s = cfg['store']
    r = 1
    span(ws, r, 1, COLS, '入会申込書 ・ 同意書', TITLE, align='center', height=26)
    r += 1
    span(ws, r, 1, COLS, f"{s['name']}　{s['address']}　TEL {s['tel']}", SUB, align='center')
    r += 1
    span(ws, r, 1, COLS, '※ご記入いただいた内容は会員管理と安全管理のために使用します。'
                         '黒のボールペンでご記入ください。', TINY, align='center')
    r += 2

    r = band(ws, r, '１．ご本人さまについて')
    r = pair(ws, r, 'お名前', 'フリガナ')
    r = pair(ws, r, '生年月日', '性別')
    r = field(ws, r, 'ご住所')
    r = pair(ws, r, '電話番号', 'メール')
    r = pair(ws, r, 'ご職業', '入会日')
    r += 1

    r = band(ws, r, '２．緊急連絡先（ご本人以外）')
    r = pair(ws, r, 'お名前', 'ご関係')
    r = pair(ws, r, '電話番号', '続柄・備考')
    r += 1

    r = band(ws, r, '３．健康状態の申告（当てはまるものに✓。1つでも✓がある場合は詳細をご記入ください）')
    r = checkboxes(ws, r, cfg['health_items'], per_row=2)
    r = field(ws, r, '✓をつけた項目の詳細', height=38)
    r = field(ws, r, '医師の許可', '☐ 医師の許可を得ている　　☐ 医師の指示はない　　☐ 該当しない', height=22)
    r = para(ws, r, '※申告もれがあると、安全なトレーニングのご提案ができません。'
                    '内容に変更があったときは、その都度お知らせください。', TINY)
    r += 1

    r = band(ws, r, '４．ご契約内容')
    span(ws, r, 1, 3, 'コース', HEAD, NAVY, BOX, align='center')
    span(ws, r, 4, 5, '料金（税込）', HEAD, NAVY, BOX, align='center')
    span(ws, r, 6, COLS, '回数・有効期限', HEAD, NAVY, BOX, align='center')
    r += 1
    for name, price, term in cfg['plans']:
        span(ws, r, 1, 3, CHECK + name, SMALL, None, BOX)
        span(ws, r, 4, 5, price, SMALL, None, BOX, align='center')
        span(ws, r, 6, COLS, term, SMALL, None, BOX)
        ws.row_dimensions[r].height = 19
        r += 1
    r = field(ws, r, 'お支払い方法', '　'.join(CHECK + m for m in cfg['payment_methods']), height=22)
    r = pair(ws, r, '初回お支払い額', '毎月のお支払日')
    r = para(ws, r, '※契約期間・料金・中途解約時の取り扱いは、別紙の契約書面（重要事項の説明を含む）'
                    'のとおりです。ご不明な点は署名前にお尋ねください。', TINY)
    r += 1

    r = band(ws, r, '５．同意事項（お読みいただき、ご同意のうえ署名をお願いします）')
    for i, item in enumerate(cfg['consent_items'], 1):
        span(ws, r, 1, 1, f'{i}.', SMALL, None, None, align='center', valign='top')
        lines = max(1, -(-len(item) // 52))
        span(ws, r, 2, COLS, item, SMALL, None, None, wrap=True, align='left',
             valign='top', height=13.5 * lines + 3)
        r += 1
    r += 1

    r = band(ws, r, '６．写真・動画の掲載について（任意）')
    r = para(ws, r, cfg['photo_consent'], SMALL)
    r = field(ws, r, '掲載の可否', '☐ 同意する　　　☐ 同意しない', height=22)
    r += 1

    r = band(ws, r, '７．署名')
    r = para(ws, r, '上記の内容を確認し、同意のうえ入会を申し込みます。', BODY)
    span(ws, r, 1, 2, '日付', BOLD, LABEL, BOX)
    span(ws, r, 3, 4, '　　　年　　月　　日', SMALL, None, BOX, align='center')
    span(ws, r, 5, 5, 'ご署名', BOLD, LABEL, BOX)
    span(ws, r, 6, COLS, '', BODY, SIGN, BOX)
    ws.row_dimensions[r].height = 34
    r += 1
    span(ws, r, 1, 2, '親権者署名', BOLD, LABEL, BOX)
    span(ws, r, 3, COLS, '（18歳未満の方は親権者のご署名をお願いします）', TINY, SIGN, BOX)
    ws.row_dimensions[r].height = 34
    r += 2

    r = band(ws, r, '【店舗記入欄】')
    r = pair(ws, r, '受付者', '受付日')
    r = field(ws, r, '確認チェック',
              '☐ 健康申告の確認　☐ 医師の許可の確認　☐ 契約書面を交付　☐ 署名もれなし　'
              '☐ 決済登録　☐ 受付台帳に記録　☐ 書類を施錠保管', height=34)
    r = field(ws, r, '会員番号', height=22)
    ws.print_area = f'A1:H{r}'
    return ws


# ------------------------------------------------------------
#  ② 休会届
# ------------------------------------------------------------
def sheet_leave(wb, cfg):
    ws = wb.create_sheet('②休会届')
    setup(ws)
    s = cfg['store']
    lv = cfg['leave']
    r = 1
    span(ws, r, 1, COLS, '休 会 届', TITLE, align='center', height=26)
    r += 1
    span(ws, r, 1, COLS, f"{s['name']}　TEL {s['tel']}", SUB, align='center')
    r += 2

    r = band(ws, r, '１．ご本人さまについて')
    r = pair(ws, r, '会員番号', 'お名前')
    r = pair(ws, r, '電話番号', '提出日')
    r += 1

    r = band(ws, r, '２．休会の期間')
    r = pair(ws, r, '休会の開始月', '再開の予定月')
    r = field(ws, r, '最終利用日', height=22)
    r += 1

    r = band(ws, r, '３．休会の理由（当てはまるものに✓）')
    r = checkboxes(ws, r, cfg['leave_reasons'], per_row=2)
    r = field(ws, r, 'ご自由記入', height=34)
    r += 1

    r = band(ws, r, '４．休会についてのご確認')
    for text in [cfg['deadlines']['leave'], lv['fee'], lv['max'], lv['restart'],
                 cfg['deadlines']['note'],
                 '休会中は予約・ご利用ができません。再開のご連絡がない場合、'
                 '休会期間の満了後は自動的に通常の会費に戻ります（要確認）。']:
        span(ws, r, 1, 1, CHECK, SMALL, None, None, align='center')
        lines = max(1, -(-len(text) // 54))
        span(ws, r, 2, COLS, text, SMALL, None, None, wrap=True, valign='top',
             height=13.5 * lines + 3)
        r += 1
    r += 1

    r = band(ws, r, '５．署名')
    r = para(ws, r, '上記の内容を確認し、休会を申し出ます。', BODY)
    span(ws, r, 1, 2, '日付', BOLD, LABEL, BOX)
    span(ws, r, 3, 4, '　　　年　　月　　日', SMALL, None, BOX, align='center')
    span(ws, r, 5, 5, 'ご署名', BOLD, LABEL, BOX)
    span(ws, r, 6, COLS, '', BODY, SIGN, BOX)
    ws.row_dimensions[r].height = 34
    r += 2

    r = band(ws, r, '【店舗記入欄】')
    r = pair(ws, r, '受付者', '受付日')
    r = field(ws, r, '処理チェック',
              '☐ 締切内に受付　☐ 休会費に変更　☐ 引き落とし金額の変更　'
              '☐ 再開予定月を台帳に記録　☐ 控えをお渡し　☐ 予約枠の解放', height=34)
    r = field(ws, r, '備考', height=26)
    ws.print_area = f'A1:H{r}'
    return ws


# ------------------------------------------------------------
#  ③ 退会届
# ------------------------------------------------------------
def sheet_withdraw(wb, cfg):
    ws = wb.create_sheet('③退会届')
    setup(ws)
    s = cfg['store']
    r = 1
    span(ws, r, 1, COLS, '退 会 届', TITLE, align='center', height=26)
    r += 1
    span(ws, r, 1, COLS, f"{s['name']}　TEL {s['tel']}", SUB, align='center')
    r += 2

    r = band(ws, r, '１．ご本人さまについて')
    r = pair(ws, r, '会員番号', 'お名前')
    r = pair(ws, r, '電話番号', '提出日')
    r += 1

    r = band(ws, r, '２．退会の希望日')
    r = pair(ws, r, '退会希望日', '最終利用日')
    r += 1

    r = band(ws, r, '３．退会の理由（当てはまるものに✓。今後の改善に使わせていただきます）')
    r = checkboxes(ws, r, cfg['withdraw_reasons'], per_row=3)
    r = field(ws, r, 'ご自由記入', height=38)
    r = field(ws, r, 'また通えるとしたら', '☐ 条件が合えば再開したい　　☐ 未定　　☐ 再開の予定はない',
              height=22)
    r += 1

    r = band(ws, r, '４．退会についてのご確認')
    for text in [cfg['deadlines']['withdraw'], cfg['deadlines']['note'],
                 '退会日をもって予約・ご利用ができなくなります。',
                 'ロッカーの鍵・貸与品は退会日までにご返却ください。',
                 '会費のお支払いが残っている場合は、退会日までに精算をお願いします。',
                 '契約期間の途中で解約される場合の取り扱いは、別紙の契約書面のとおりです'
                 '（法律上の中途解約・クーリング・オフの権利は制限されません）。']:
        span(ws, r, 1, 1, CHECK, SMALL, None, None, align='center')
        lines = max(1, -(-len(text) // 54))
        span(ws, r, 2, COLS, text, SMALL, None, None, wrap=True, valign='top',
             height=13.5 * lines + 3)
        r += 1
    r += 1

    r = band(ws, r, '５．署名')
    r = para(ws, r, '上記の内容を確認し、退会を申し出ます。', BODY)
    span(ws, r, 1, 2, '日付', BOLD, LABEL, BOX)
    span(ws, r, 3, 4, '　　　年　　月　　日', SMALL, None, BOX, align='center')
    span(ws, r, 5, 5, 'ご署名', BOLD, LABEL, BOX)
    span(ws, r, 6, COLS, '', BODY, SIGN, BOX)
    ws.row_dimensions[r].height = 34
    r += 2

    r = band(ws, r, '【店舗記入欄】')
    r = pair(ws, r, '受付者', '受付日')
    r = field(ws, r, '処理チェック',
              '☐ 締切内に受付　☐ 引き落とし停止　☐ 貸与物の返却確認　☐ 未収金の精算　'
              '☐ 退会理由を記録　☐ 受付台帳を締め　☐ 会員データを退会に変更', height=34)
    r = field(ws, r, '引き止めの打ち手（実施したこと）', height=30)
    ws.print_area = f'A1:H{r}'
    return ws


# ------------------------------------------------------------
#  ④ 受付台帳（処理もれを出さないための一覧）
# ------------------------------------------------------------
def sheet_ledger(wb, cfg):
    ws = wb.create_sheet('④受付台帳')
    ws.sheet_view.showGridLines = False
    ws['A1'] = '手続き受付台帳（入会・休会・再開・退会・プラン変更）'
    ws['A1'].font = TITLE
    ws['A2'] = '書類を受け取ったら、その日のうちにこの台帳へ記録する。処理が終わったら「処理済」にする。'
    ws['A2'].font = SUB

    heads = ['受付日', '会員番号', 'お名前', '種別', '適用月', '締切内',
             '書類の署名', '決済の変更', '控えを渡した', '処理状況', '受付者', '備考']
    widths = [11, 10, 14, 11, 10, 8, 10, 10, 11, 10, 9, 26]
    HROW = 4
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=HROW, column=i, value=h)
        c.font = HEAD
        c.fill = NAVY
        c.border = BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    first = HROW + 1
    ROWS = 40
    last = first + ROWS - 1
    for r in range(first, last + 1):
        for i in range(1, len(heads) + 1):
            c = ws.cell(row=r, column=i)
            c.border = Border(left=hair, right=hair, top=hair, bottom=hair)
            c.font = BODY
            c.alignment = Alignment(horizontal='center' if i != 12 else 'left')
        ws.cell(row=r, column=1).number_format = 'yyyy/m/d'
        ws.cell(row=r, column=5).number_format = 'yyyy/m'

    dv_type = DataValidation(type='list', formula1='"' + ','.join(cfg['ledger_types']) + '"',
                            allow_blank=True, showErrorMessage=True)
    dv_yn = DataValidation(type='list', formula1='"✓,ー"', allow_blank=True)
    dv_st = DataValidation(type='list', formula1='"未処理,処理中,処理済"', allow_blank=True)
    for dv, rng in [(dv_type, f'D{first}:D{last}'), (dv_yn, f'F{first}:I{last}'),
                    (dv_st, f'J{first}:J{last}')]:
        ws.add_data_validation(dv)
        dv.add(rng)

    # 未処理が赤く残る＝締切前に気づける
    ws.conditional_formatting.add(
        f'J{first}:J{last}',
        FormulaRule(formula=[f'$J{first}="未処理"'],
                    fill=PatternFill('solid', bgColor='FFC7CE'), font=Font(color='9C0006')))
    ws.conditional_formatting.add(
        f'J{first}:J{last}',
        FormulaRule(formula=[f'$J{first}="処理済"'],
                    fill=PatternFill('solid', bgColor='C6EFCE'), font=Font(color='006100')))

    r = last + 2
    ws.cell(row=r, column=1, value='件数（当月の集計に使う）').font = H2
    r += 1
    labels = cfg['ledger_types'] + ['未処理']
    for i, label in enumerate(labels):
        col = 1 + i * 2
        span(ws, r, col, col + 1, label, HEAD, NAVY, BOX, align='center')
        target = 'J' if label == '未処理' else 'D'
        span(ws, r + 1, col, col + 1,
             f'=COUNTIF({target}{first}:{target}{last},"{label}")', BOLD, None, BOX,
             align='center')
    ws.freeze_panes = f'A{first}'
    setup_print(ws, landscape=True, title_rows=f'{HROW}:{HROW}')
    return ws


# ------------------------------------------------------------
#  ⑤ 手続きルール
# ------------------------------------------------------------
def sheet_rules(wb, cfg):
    ws = wb.create_sheet('⑤手続きルール')
    ws.sheet_view.showGridLines = False
    ws['A1'] = '手続きの締切と流れ'
    ws['A1'].font = TITLE
    ws['A2'] = '締切を1日過ぎると、翌月分の会費が発生してトラブルになる。締切前の金曜に必ず取りまとめる。'
    ws['A2'].font = SUB
    heads = ['種別', '締切', '会員さまにお願いすること', '店舗がやること', '効いてくる数字']
    widths = [12, 22, 26, 46, 18]
    HROW = 4
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=HROW, column=i, value=h)
        c.font = HEAD
        c.fill = NAVY
        c.border = BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    r = HROW
    for row in cfg['rules_flow']:
        r += 1
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = BOLD if i == 1 else BODY
            c.border = BOX
            if i == 1:
                c.fill = LABEL
            c.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[r].height = 46
    r += 2
    ws.cell(row=r, column=1, value='書類の保管').font = H2
    for text in ['署名のある書類は原本を施錠した場所に保管する（健康情報は要配慮個人情報）。',
                 '会員さまには必ず控えをお渡しする（休会・退会は控えがないと後で揉める）。',
                 'スマホで撮った画像を個人の端末に残さない。',
                 '退会した方の書類も、法令・契約上の保存期間が終わるまでは廃棄しない（保存期間は要確認）。']:
        r += 1
        ws.cell(row=r, column=1, value='・').alignment = Alignment(horizontal='center')
        c = ws.cell(row=r, column=2, value=text)
        c.font = SMALL
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 18
    setup_print(ws, landscape=True, one_page=True)
    return ws


# ------------------------------------------------------------
#  ⑥ 法務チェック
# ------------------------------------------------------------
def sheet_legal(wb, cfg):
    ws = wb.create_sheet('⑥法務チェック')
    ws.sheet_view.showGridLines = False
    ws['A1'] = '配布する前に確認すること'
    ws['A1'].font = TITLE
    ws['A2'] = 'この書式はひな形です。下の項目を埋め、最終文面は専門家（弁護士等）の確認を受けてください。'
    ws['A2'].font = RED
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 100
    r = 4
    for i, text in enumerate(cfg['legal_notes'], 1):
        c = ws.cell(row=r, column=1, value=CHECK)
        c.alignment = Alignment(horizontal='center', vertical='top')
        c2 = ws.cell(row=r, column=2, value=text)
        c2.font = BODY
        c2.alignment = Alignment(wrap_text=True, vertical='top')
        c2.fill = WARN
        c2.border = BOX
        ws.row_dimensions[r].height = 15 * max(1, -(-len(text) // 58)) + 6
        r += 2
    setup_print(ws, one_page=True)
    return ws


def main():
    cfg_path = arg('config', os.path.join(ROOT, 'config', 'member-forms.json'))
    cfg = json.load(open(cfg_path, encoding='utf-8'))
    if arg('store'):
        cfg['store']['name'] = arg('store')

    wb = Workbook()
    wb.remove(wb.active)
    sheet_join(wb, cfg)
    sheet_leave(wb, cfg)
    sheet_withdraw(wb, cfg)
    sheet_ledger(wb, cfg)
    sheet_rules(wb, cfg)
    sheet_legal(wb, cfg)

    name = cfg['store']['name'].replace('/', '／')
    out = arg('out', os.path.join(ROOT, 'out', f'会員手続き書類_{name}.xlsx'))
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f'できました: {out}')
    print(f"  シート: {' / '.join(ws.title for ws in wb.worksheets)}")
    print('  ①〜③はA4たてで印刷してそのまま記入できます。')
    print('  ★配布前に「⑥法務チェック」の項目（◯・要確認の箇所）を必ず埋めてください。')


if __name__ == '__main__':
    main()
