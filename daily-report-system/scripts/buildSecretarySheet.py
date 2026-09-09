#!/usr/bin/env python3
# ============================================================
#  AI秘書 1日最適化シート（Excel）の生成
# ------------------------------------------------------------
#  毎朝のプランを、印刷しても使える1冊の Excel にします。
#
#  実行:
#    node scripts/secretary.js --file=today.txt --json | \
#      python3 scripts/buildSecretarySheet.py
#    （または npm run secretary:sheet -- --file=today.txt）
#
#  シート構成:
#    ① 今日の作戦   … 絶対終わらせる3つ・数字・秘書からの指摘
#    ② タイムライン … 時刻ごとの動き（チェック欄つき）
#    ③ 依頼リスト   … 誰に何を頼むか（そのまま指示書になる）
#    ④ 相手待ち     … 自分の作業ではないが、催促の判断がいるもの
#    ⑤ 今日やらない … 落とした理由と、いつやるか
# ============================================================
import json
import os
import sys
from datetime import date as _date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

FONT = 'Yu Gothic'
NAVY = '1F3864'
INK = '262626'
GRAY = '767676'

H1 = Font(name=FONT, size=16, bold=True, color=NAVY)
H2 = Font(name=FONT, size=11, bold=True, color='FFFFFF')
BODY = Font(name=FONT, size=10, color=INK)
BOLD = Font(name=FONT, size=10, bold=True, color=INK)
SMALL = Font(name=FONT, size=9, color=GRAY)
BIG = Font(name=FONT, size=12, bold=True, color=INK)

HEAD_FILL = PatternFill('solid', fgColor=NAVY)
ZEBRA = PatternFill('solid', fgColor='F7F9FC')
MUST_FILL = PatternFill('solid', fgColor='FFF3CD')
FIXED_FILL = PatternFill('solid', fgColor='F0F0F0')

THIN = Side(style='thin', color='D0D7E5')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical='center', wrap_text=True)
MID = Alignment(vertical='center')
CENTER = Alignment(horizontal='center', vertical='center')

# カテゴリごとの色（画面の絵文字と対応させ、紙でも見分けられるようにする）
CAT = {
    'THINK':      ('🧠 THINK',      'リサーチ・思考',   'E8EEF9', '2F5597'),
    'IMPROVE':    ('🛠 IMPROVE',    '改善・制作',       'E7F2E9', '375623'),
    'DELEGATE':   ('📤 DELEGATE',   '依頼・指示',       'FDEEE0', 'C55A11'),
    'MEETING':    ('👥 MEETING',    '会議・対話',       'F0E8F6', '7030A0'),
    'MANAGEMENT': ('📦 MANAGEMENT', '管理・確認',       'E3F0F3', '1F6F7A'),
    'ADMIN':      ('📑 ADMIN',      '事務・単純作業',   'EFEFEF', '595959'),
}
WEEK = ['月', '火', '水', '木', '金', '土', '日']


def ask_text(title, who):
    """依頼リスト用に「◯◯さんに〜を依頼」から相手と語尾を落とす（相手は別列にあるため）"""
    t = str(title)
    if who:
        for name in (who, who.rstrip('さん')):
            if name and t.startswith(name):
                t = t[len(name):].lstrip('にへ　 ')
                break
    for suf in ('を依頼', 'の依頼', 'を進める', 'を連携', 'と連携', 'を催促'):
        if t.endswith(suf):
            t = t[: -len(suf)]
            break
    # 「…の算出をこころさん」のように相手が末尾に残る形も落とす
    if who:
        for name in (who, who.rstrip('さん')):
            if name and t.endswith(name):
                t = t[: -len(name)]
                break
    t = t.rstrip('をにへと　 ')
    return t.strip() or str(title)


def hm(minutes):
    m = int(minutes or 0)
    if m < 60:
        return f'{m}分'
    return f'{m // 60}時間{m % 60}分' if m % 60 else f'{m // 60}時間'


def setup(ws, widths, title):
    """見出し行とページ設定をまとめて行う"""
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:3'
    ws['A1'] = title
    ws['A1'].font = H1
    ws.row_dimensions[1].height = 26


def header(ws, row, labels):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = H2
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BOX
    ws.row_dimensions[row].height = 22


def put(ws, row, values, fill=None, bold=False, heights=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = BOLD if bold else BODY
        c.alignment = WRAP
        c.border = BOX
        if fill:
            c.fill = fill
    if heights:
        ws.row_dimensions[row].height = heights


def sheet_strategy(wb, p):
    ws = wb.create_sheet('① 今日の作戦')
    setup(ws, [4, 34, 16, 62, 14], f"AI秘書 1日最適化　{p['date']}（{WEEK[_date.fromisoformat(p['date']).weekday()]}）")

    ws['A2'] = (f"タスク {p['tasks_count']}件 ／ 見積合計 {hm(p['total_minutes'])} ／ "
                f"今日つかえる時間 {hm(p['capacity'])} ／ 配置済み {hm(p['planned_minutes'])}")
    ws['A2'].font = SMALL

    r = 4
    ws.cell(row=r, column=1, value='🔥 今日絶対終わらせる3つ').font = BIG
    r += 1
    header(ws, r, ['#', 'タスク', '時間帯', 'なぜ今日やるのか', '完了'])
    for i, t in enumerate(p['must_do'], start=1):
        r += 1
        put(ws, r, [i, f"{t['icon']} {t['title']}", t['slot'] or '-', t['why'], '□'],
            fill=MUST_FILL, heights=34)
        ws.cell(row=r, column=1).alignment = CENTER
        ws.cell(row=r, column=5).alignment = CENTER

    r += 2
    ws.cell(row=r, column=1, value='⚠️ 秘書からの指摘').font = BIG
    r += 1
    header(ws, r, ['#', '指摘', '', '', ''])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    for i, a in enumerate(p['advice'], start=1):
        r += 1
        put(ws, r, [i, a], fill=ZEBRA if i % 2 else None, heights=30)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.cell(row=r, column=1).alignment = CENTER
    return ws


def sheet_timeline(wb, p):
    ws = wb.create_sheet('② タイムライン')
    setup(ws, [14, 15, 52, 10, 22, 8], f"今日のスケジュール　{p['date']}")
    ws['A2'] = '★＝今日絶対終わらせる3つ　📌＝時間が決まっている予定　網掛け＝食事・休憩・固定枠'
    ws['A2'].font = SMALL
    header(ws, 3, ['時刻', '区分', '内容', '所要', '相手／メモ', '完了'])
    r = 3
    for e in p['timeline']:
        r += 1
        kind = e['kind']
        if kind == 'task':
            name, _, bg, fg = CAT[e['category']]
            mark = '★ ' if e.get('must') else ''
            pin = '📌 ' if e.get('pinned') else ''
            part = f"［{e['part']}］" if e.get('part') else ''
            # 相手を出すのは「人に振る仕事」だけ。自分の作業に担当者名が出ると誤解を生む。
            memo = (e.get('who') or '') if e['category'] == 'DELEGATE' else ''
            put(ws, r, [f"{e['start']}〜{e['end']}", name, f"{mark}{pin}{e['title']}{part}",
                        hm(e['minutes']), memo, '□'],
                fill=PatternFill('solid', fgColor=bg), heights=24)
            ws.cell(row=r, column=2).font = Font(name=FONT, size=9, bold=True, color=fg)
            if e.get('must'):
                ws.cell(row=r, column=3).font = Font(name=FONT, size=10, bold=True, color=INK)
        else:
            label = {'break': '休憩', 'free': '空き', 'fixed': '固定'}.get(kind, kind)
            put(ws, r, [f"{e['start']}〜{e['end']}", label, e.get('label') or '',
                        hm(e['minutes']), '', ''], fill=FIXED_FILL, heights=20)
            for col in range(1, 7):
                ws.cell(row=r, column=col).font = SMALL
        ws.cell(row=r, column=1).alignment = CENTER
        ws.cell(row=r, column=4).alignment = CENTER
        ws.cell(row=r, column=6).alignment = CENTER
    ws.freeze_panes = 'A4'
    return ws


def sheet_delegate(wb, p):
    ws = wb.create_sheet('③ 依頼リスト')
    setup(ws, [6, 22, 54, 12, 16, 8], '朝一で人に振る仕事（このまま指示書として使えます）')
    ws['A2'] = '他人の着手待ちを最初に解くほど、1日の総処理量が増えます。上から順に出してください。'
    ws['A2'].font = SMALL
    header(ws, 3, ['#', '誰に', '何を依頼するか', '目安', '出す時間', '済'])
    r = 3
    for i, t in enumerate(p['morning_delegate'], start=1):
        r += 1
        who = t.get('who') or '担当者を決める'
        put(ws, r, [i, who + ('' if t.get('named') else '（推奨）'), ask_text(t['title'], t.get('who')),
                    hm(t['minutes']), t.get('slot') or '', '□'],
            fill=ZEBRA if i % 2 else None, heights=26)
        ws.cell(row=r, column=1).alignment = CENTER
        ws.cell(row=r, column=2).font = BOLD
        for col in (4, 5, 6):
            ws.cell(row=r, column=col).alignment = CENTER
    r += 2
    ws.cell(row=r, column=1, value='📋 カテゴリ別のタスク仕分け').font = BIG
    r += 1
    header(ws, r, ['#', '区分', 'タスク', '所要', '時間帯', ''])
    n = 0
    for key, b in p['buckets'].items():
        name, _, bg, fg = CAT[key]
        for t in b['items']:
            n += 1
            r += 1
            put(ws, r, [n, name, t['title'], hm(t['minutes']), t.get('slot') or '', ''],
                fill=PatternFill('solid', fgColor=bg), heights=22)
            ws.cell(row=r, column=1).alignment = CENTER
            ws.cell(row=r, column=2).font = Font(name=FONT, size=9, bold=True, color=fg)
    return ws


def sheet_waiting(wb, p):
    ws = wb.create_sheet('④ 相手待ち')
    setup(ws, [6, 56, 20, 14, 24], '相手待ち（自分の作業ではありません。催促するかどうかだけ決めます）')
    ws['A2'] = '3日以上たったものは、催促するか「待たずに進める方法」に切り替えてください。'
    ws['A2'].font = SMALL
    header(ws, 3, ['#', '内容', '待っている相手', '経過', '判断'])
    r = 3
    for i, w in enumerate(p['waiting'], start=1):
        r += 1
        days = w.get('days', 0)
        put(ws, r, [i, w['title'], w.get('who') or '相手',
                    '本日から' if days == 0 else f'{days}日経過',
                    '催促する' if days >= 3 else ''],
            fill=PatternFill('solid', fgColor='FDECEA') if days >= 3 else (ZEBRA if i % 2 else None),
            heights=24)
        ws.cell(row=r, column=1).alignment = CENTER
        for col in (3, 4, 5):
            ws.cell(row=r, column=col).alignment = CENTER
    if not p['waiting']:
        put(ws, 4, ['', '相手待ちはありません', '', '', ''])
    return ws


def sheet_later(wb, p):
    ws = wb.create_sheet('⑤ 今日やらない')
    setup(ws, [6, 46, 16, 44, 22], '今日やらなくていい仕事（切ったものと、その理由）')
    ws['A2'] = '全部を今日やろうとすると全部が中途半端になります。ここは「意図して落とした」ものです。'
    ws['A2'].font = SMALL
    header(ws, 3, ['#', 'タスク', '所要', '後回しにしてよい理由', '推奨（翌営業日）'])
    r = 3
    for i, t in enumerate(p['deferred'], start=1):
        r += 1
        put(ws, r, [i, f"{t['icon']} {t['title']}", hm(t['minutes']), t['reason'], t['recommend']],
            fill=ZEBRA if i % 2 else None, heights=28)
        ws.cell(row=r, column=1).alignment = CENTER
        ws.cell(row=r, column=3).alignment = CENTER
    if not p['deferred']:
        put(ws, 4, ['', '今日のタスクはすべて今日の枠に収まります', '', '', ''])
    return ws


def main():
    raw = sys.stdin.read() if not sys.stdin.isatty() else None
    if not raw:
        src = next((a.split('=', 1)[1] for a in sys.argv[1:] if a.startswith('--in=')), None)
        if not src:
            print('プランのJSONを渡してください（secretary.js --json の出力）', file=sys.stderr)
            return 1
        raw = open(src, encoding='utf-8').read()
    p = json.loads(raw)

    wb = Workbook()
    wb.remove(wb.active)
    sheet_strategy(wb, p)
    sheet_timeline(wb, p)
    sheet_delegate(wb, p)
    sheet_waiting(wb, p)
    sheet_later(wb, p)
    wb.properties.title = f"AI秘書 1日最適化 {p['date']}"

    out_dir = next((a.split('=', 1)[1] for a in sys.argv[1:] if a.startswith('--out-dir=')), os.path.join(ROOT, 'out'))
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, f"AI秘書_{p['date']}.xlsx")
    wb.save(out)
    print(out)
    return 0


if __name__ == '__main__':
    sys.exit(main())
