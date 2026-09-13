#!/usr/bin/env python3
# ============================================================
#  韓国SNS 案件依頼 管理表（Excel・1枚もの）
# ------------------------------------------------------------
#  インフルエンサーへの「案件依頼」だけを追うシンプルな表です。
#  成約 → 来店 → 投稿 の3ステップだけを見ます。
#  （撮影・編集・広告・予約まで追うのは別表 buildKoreaSnsSheet.py）
#
#  列: 成約日 / 成約担当者 / アカウント名 / 案件依頼費用 / フォロワー数 /
#      来店日 / 投稿予定日 / 投稿形式 / 状況 ＋（任意）店舗・投稿URL
#
#  自動でやること:
#    ・上の段に 件数・費用合計・状況別の件数・平均フォロワー数
#    ・来店日を過ぎても「来店待ち」の行を赤くする
#    ・投稿予定日を過ぎても「投稿済み」でない行を赤くする
#
#  実行:
#    python3 scripts/buildKoreaOrderSheet.py
#    python3 scripts/buildKoreaOrderSheet.py --month=2026-09
#    python3 scripts/buildKoreaOrderSheet.py --out=/path/to/file.xlsx
# ============================================================
import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG = os.path.join(ROOT, 'config', 'korea-sns.json')

F = 'Yu Gothic'
NUM = '#,##0'
YEN = '¥#,##0;(¥#,##0);-'
DATE_FMT = 'yyyy/mm/dd'

HDR_FILL = PatternFill('solid', fgColor='188038')
INPUT_FILL = PatternFill('solid', fgColor='FFF7CC')
SUM_FILL = PatternFill('solid', fgColor='E2EFDA')
NG_FILL = PatternFill('solid', fgColor='FCE4E4')

HDR_FONT = Font(name=F, bold=True, color='FFFFFF', size=10)
BOLD = Font(name=F, bold=True)
NORM = Font(name=F, size=10)
SAMPLE = Font(name=F, size=10, italic=True, color='808080')
NOTE = Font(name=F, size=9, color='808080')
TITLE = Font(name=F, bold=True, size=14)
THIN = Border(*[Side(style='thin', color='BFBFBF')] * 4)

HEAD_ROW = 6       # 見出しの行（上の5行はサマリー）
FIRST_ROW = 7      # 記入例の行（＝データの1行目）
LAST_ROW = 206     # 200行ぶん用意する

FORMATS = ['リール', 'ストーリー', 'フィード']
STATUS = ['来店待ち', '投稿待ち', '投稿済み']

# 列（ご指定の順番のまま。末尾2つは任意）
COLS = [
    ('成約日', 12), ('成約担当者', 14), ('アカウント名', 22), ('案件依頼費用', 14),
    ('フォロワー数', 13), ('来店日', 12), ('投稿予定日', 12), ('投稿形式', 12), ('状況', 12),
    ('店舗（任意）', 12), ('投稿URL（任意）', 34),
]
C_DATE, C_OWNER, C_ACC, C_COST, C_FOL, C_VISIT, C_PLAN, C_FORM, C_STATUS, C_SHOP, C_URL = range(1, 12)


def arg(name, default=None):
    for a in sys.argv[1:]:
        if a.startswith(f'--{name}='):
            return a[len(name) + 3:]
    return default


def inline_list(values):
    """プルダウンをその場で指定する（別シートを作らないぶん、表が1枚で済む）。"""
    return '"' + ','.join(values) + '"'


def build(wb, shops):
    ws = wb.active
    ws.title = '案件依頼'

    ws['A1'] = '韓国SNS 案件依頼 管理表'
    ws['A1'].font = TITLE
    ws['A2'] = '成約したインフルエンサー案件を1件1行で。黄色の欄を埋めるだけです。'
    ws['A2'].font = NOTE
    ws['A3'] = '来店日・投稿予定日を過ぎても状況が進んでいない行は、自動で赤くなります。'
    ws['A3'].font = NOTE

    # ── 上段のサマリー（数式）──
    col = get_column_letter
    rng = lambda c: f'${col(c)}${FIRST_ROW}:${col(c)}${LAST_ROW}'  # noqa: E731
    summary = [
        ('件数', f'=COUNTA({rng(C_DATE)})'),
        ('依頼費用の合計', f'=SUM({rng(C_COST)})', YEN),
        ('平均フォロワー', f'=IF(COUNT({rng(C_FOL)})=0,"",ROUND(AVERAGE({rng(C_FOL)}),0))'),
        ('来店待ち', f'=COUNTIF({rng(C_STATUS)},"{STATUS[0]}")'),
        ('投稿待ち', f'=COUNTIF({rng(C_STATUS)},"{STATUS[1]}")'),
        ('投稿済み', f'=COUNTIF({rng(C_STATUS)},"{STATUS[2]}")'),
    ]
    for i, item in enumerate(summary):
        label, formula = item[0], item[1]
        fmt = item[2] if len(item) > 2 else NUM
        c1 = ws.cell(row=4, column=1 + i * 2, value=label)
        c1.font = BOLD
        c1.fill = SUM_FILL
        c1.border = THIN
        c1.alignment = Alignment(horizontal='center')
        c2 = ws.cell(row=5, column=1 + i * 2, value=formula)
        c2.font = BOLD
        c2.fill = SUM_FILL
        c2.border = THIN
        c2.number_format = fmt
        c2.alignment = Alignment(horizontal='center')

    # ── 見出し ──
    for i, (label, width) in enumerate(COLS, start=1):
        c = ws.cell(row=HEAD_ROW, column=i, value=label)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, THIN
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[col(i)].width = width
    ws.row_dimensions[HEAD_ROW].height = 24
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=3)

    # ── 入力欄 ──
    for r in range(FIRST_ROW, LAST_ROW + 1):
        for i in range(1, len(COLS) + 1):
            c = ws.cell(row=r, column=i)
            c.font = SAMPLE if r == FIRST_ROW else NORM
            c.border = THIN
            c.fill = INPUT_FILL
            c.alignment = Alignment(vertical='center')
        for i in (C_DATE, C_VISIT, C_PLAN):
            ws.cell(row=r, column=i).number_format = DATE_FMT
        ws.cell(row=r, column=C_COST).number_format = YEN
        ws.cell(row=r, column=C_FOL).number_format = NUM

    # ── プルダウン ──
    for values, c in ((FORMATS, C_FORM), (STATUS, C_STATUS), (shops, C_SHOP)):
        v = DataValidation(type='list', formula1=inline_list(values), allow_blank=True, showDropDown=False)
        ws.add_data_validation(v)
        v.add(f'{col(c)}{FIRST_ROW}:{col(c)}{LAST_ROW}')

    # ── 遅れている行を赤くする ──
    last_col = col(len(COLS))
    body = f'A{FIRST_ROW}:{last_col}{LAST_ROW}'
    # 来店日を過ぎたのに「来店待ち」のまま
    ws.conditional_formatting.add(body, FormulaRule(
        formula=[f'AND($A{FIRST_ROW}<>"",$I{FIRST_ROW}="{STATUS[0]}",$F{FIRST_ROW}<>"",'
                 f'$F{FIRST_ROW}<TODAY())'],
        fill=NG_FILL, stopIfTrue=False))
    # 投稿予定日を過ぎたのに「投稿済み」になっていない
    ws.conditional_formatting.add(body, FormulaRule(
        formula=[f'AND($A{FIRST_ROW}<>"",$I{FIRST_ROW}<>"{STATUS[2]}",$G{FIRST_ROW}<>"",'
                 f'$G{FIRST_ROW}<TODAY())'],
        fill=NG_FILL, stopIfTrue=False))

    # ── 記入例 ──
    sample = [date(2026, 9, 12), '北野', '@myeongdong_kim', 150000, 52000,
              date(2026, 9, 20), date(2026, 9, 25), 'リール', '来店待ち',
              shops[0], 'https://www.instagram.com/reel/xxxxxxxx/']
    for i, v in enumerate(sample, start=1):
        ws.cell(row=FIRST_ROW, column=i, value=v).font = SAMPLE
    ws.cell(row=FIRST_ROW, column=len(COLS) + 1, value='← 記入例。使うときは消してください').font = NOTE

    r = LAST_ROW + 2
    for k, line in enumerate([
        '状況の進め方: 成約したら「来店待ち」→ 来店したら「投稿待ち」→ 投稿されたら「投稿済み」。',
        '投稿URLは任意ですが、入れておくと「いつ・どんな形で投稿されたか」があとから確認できます。',
    ]):
        ws.cell(row=r + k, column=1, value=line).font = NOTE
    return ws


def main():
    with open(CONFIG, encoding='utf-8') as f:
        cfg = json.load(f)
    shops = [a['id'] for a in cfg['accounts']]

    month = arg('month', date.today().strftime('%Y-%m'))
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True  # 開いた瞬間に集計を計算させる
    build(wb, shops)

    out = arg('out') or os.path.join(ROOT, 'out', f'韓国SNS_案件依頼管理_{month}.xlsx')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f'作成しました: {out}')
    print('  （1行目は記入例です。使うときは消してください）')


if __name__ == '__main__':
    main()
