#!/usr/bin/env python3
"""在庫消化カレンダー（発注前にいつも作る形式）を生成する。

  python3 scripts/buildBurndownCalendar.py
  → out/在庫消化カレンダー.xlsx （単一シート）

レイアウトは既存のGoogleスプレッドシート版に合わせている:
  1行目 タイトル / 2行目 凡例 / 3行目 前提メモ
  5行目 月見出し / 6行目 見出し（SKU（色）〜在庫月数＋日付）
  7行目以降 ■グループごとのSKU行、右側に日次カレンダー
  凡例: 赤=欠品 緑=入荷日 青=在庫あり 黄=売切間近(7日分以下) 灰=販売なし
"""
import csv, datetime, json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'data')
OUT = os.path.join(ROOT, 'out')
STOCK_DATE = datetime.date(2026, 9, 4)
CAL_END = datetime.date(2027, 6, 30)
TRANSIT = 14

thin = Side(style='thin', color='D0D0D0')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill('solid', fgColor='1F4E78')
HFo = Font(bold=True, color='FFFFFF', size=10)
TF = Font(bold=True, size=14, color='1F4E78')
SUBF = Font(size=9, color='808080')
GF = PatternFill('solid', fgColor='4472C4')
GFo = Font(bold=True, color='FFFFFF', size=11)
RED = PatternFill('solid', fgColor='FF0000')
GRN = PatternFill('solid', fgColor='00B050')
BLU = PatternFill('solid', fgColor='BDD7EE')
YEL = PatternFill('solid', fgColor='FFFF00')
GRY = PatternFill('solid', fgColor='D9D9D9')
MON_A = PatternFill('solid', fgColor='1F4E78')   # 月ヘッダー（奇数月）
MON_B = PatternFill('solid', fgColor='2E75B6')   # 月ヘッダー（偶数月）
_medium = Side(style='medium', color='1F4E78')
MONBD = Border(left=_medium, right=_medium, top=_medium, bottom=_medium)
MSTART = Border(left=_medium, right=thin, top=thin, bottom=thin)  # 月初の区切り線
C = Alignment(horizontal='center', vertical='center')
LN = Alignment(horizontal='left', vertical='center')

# 日販＝7月の3媒体（Amazon・楽天・自社）実績 ÷ 31日
DPS = {
    '106': 12.84, '101': 12.07, '107': 7.82, '108': 5.72, '111': 3.48,
    '105': 1.36, '109': 0.68, '110': 0.62, '103a': 0.35,
    '206': 7.30, '207': 5.65, '201': 4.55, '208': 4.30, '211': 3.30, '210': 0.80,
    '307': 6.08, '301': 2.28, '306': 2.17, '310': 1.52, '308': 1.07,
    'N_arumi01': 2.48, 'N_arumi02': 1.06,
    'NA_M_silver': 0.00, 'NA_M_black': 0.00,
    'arumi01': 1.61, 'arumi02': 0.58,
    'zip_S_black': 0.00, 'zip_S_silver': 0.00,
    'zip_M_black': 0.00, 'zip_M_silver': 0.00,
    'outdoorsk001': 0.00, 'outdoorsk002': 0.00,
    'L_khaki': 0.00,
    'TM_S_black': 0.00, 'TM_S_gray': 0.00, 'TM_S_silver': 0.00,
    'TM_S_white': 0.00, 'TM_S_turquoise': 0.00,
    'TM_M_black': 0.00, 'TM_M_gray': 0.00, 'TM_M_silver': 0.00, 'TM_M_white': 0.00,
    'TM_L_black': 0.00, 'TM_L_gray': 0.00, 'TM_L_silver': 0.00,
    'rental_M': 1.39,
}

# 表示順（グループ名, [(表示名, SKU)]）— スプレッドシート版の並びに合わせる
GROUPS = [
    ('■ 多機能PC S', [
        ('多機能PC S マットブラック', '106'), ('多機能PC S エナメルシルバー', '101'),
        ('多機能PC S マットシルバー', '107'), ('多機能PC S マットホワイト', '108'),
        ('多機能PC S エナメルカーキ', '111'), ('多機能PC S エナメルホワイト', '105'),
        ('多機能PC S エナメルピンク', '109'), ('多機能PC S マットグレー', '110'),
        ('多機能PC S エナメルスカイブルー', '103a')]),
    ('■ 多機能PC M', [
        ('多機能PC M マットブラック', '206'), ('多機能PC M マットシルバー', '207'),
        ('多機能PC M エナメルシルバー', '201'), ('多機能PC M マットホワイト', '208'),
        ('多機能PC M エナメルカーキ', '211'), ('多機能PC M マットグレー', '210')]),
    ('■ 多機能PC L', [
        ('多機能PC L マットシルバー', '307'), ('多機能PC L エナメルシルバー', '301'),
        ('多機能PC L マットブラック', '306'), ('多機能PC L マットグレー', '310'),
        ('多機能PC L マットホワイト', '308'), ('多機能PC L エナメルカーキ(新色)', 'L_khaki')]),
    ('■ ノーマルアルミ S（在庫表では「クラシックアルミ」）', [
        ('ノーマルアルミ S シルバー', 'N_arumi01'), ('ノーマルアルミ S ブラック', 'N_arumi02')]),
    ('■ クラシックアルミ M（ノーマルアルミM）', [
        ('クラシックアルミ M シルバー', 'NA_M_silver'), ('クラシックアルミ M ブラック', 'NA_M_black')]),
    ('■ 多機能アルミ S', [
        ('多機能アルミ S シルバー', 'arumi01'), ('多機能アルミ S ブラック', 'arumi02')]),
    ('■ ジップ S', [
        ('ジップ S シルバー', 'zip_S_silver'), ('ジップ S ブラック', 'zip_S_black')]),
    ('■ ジップ M', [
        ('ジップ M シルバー', 'zip_M_silver'), ('ジップ M ブラック', 'zip_M_black')]),
    ('■ アウトドア(スキー)', [
        ('アウトドア マットシルバー', 'outdoorsk002'), ('アウトドア マットブラック', 'outdoorsk001')]),
    ('■ サザンモデル(TM窄框款PC箱)／新商品', [
        ('サザンモデル S ブラック', 'TM_S_black'), ('サザンモデル S シルバー', 'TM_S_silver'),
        ('サザンモデル S トグレー', 'TM_S_gray'), ('サザンモデル S ホワイト', 'TM_S_white'),
        ('サザンモデル S ターコイズ', 'TM_S_turquoise'),
        ('サザンモデル M ブラック', 'TM_M_black'), ('サザンモデル M シルバー', 'TM_M_silver'),
        ('サザンモデル M トグレー', 'TM_M_gray'), ('サザンモデル M ホワイト', 'TM_M_white'),
        ('サザンモデル L ブラック', 'TM_L_black'), ('サザンモデル L シルバー', 'TM_L_silver'),
        ('サザンモデル L トグレー', 'TM_L_gray')]),
    ('■ 参考', [('（参考）レンタルM', 'rental_M')]),
]

# 9/4在庫一覧に品目が無いSKU（入庫済みだが未登録の可能性）
NOT_IN_SNAPSHOT = {'rental_M'}
# 新商品・新色（9/4時点で在庫ゼロが正しい）
NEW_ITEMS = {'L_khaki', 'TM_S_black', 'TM_S_gray', 'TM_S_silver', 'TM_S_white',
             'TM_S_turquoise', 'TM_M_black', 'TM_M_gray', 'TM_M_silver', 'TM_M_white',
             'TM_L_black', 'TM_L_gray', 'TM_L_silver'}


def arrive(ship, days=TRANSIT):
    y, m, d = [int(x) for x in ship.split('/')]
    return datetime.date(y, m, d) + datetime.timedelta(days=days)


def build_arrivals(stock, dps):
    """入荷予定（到着日→SKU→数量）。

    確定便＝パッキングリストの色別内訳。
    未出荷分＝工場の上线计划（data/production-plans.json）の品番別内訳を使い、
    到着日は「工場出荷予定日＋TRANSIT日」で置く。
    """
    a = [
        # WHSU5516465 8/27出荷・9/7着（PC多機能M 700 + ノーマルアルミM 22）
        (datetime.date(2026, 9, 7), {'201': 100, '207': 200, '210': 100, '208': 100,
                                     '206': 200, 'NA_M_silver': 22}, 'WHSU5516465'),
        # TSSU5227711 8/25出荷・9/8着（PC多機能S 1,050 + ノーマルアルミS 27）
        (datetime.date(2026, 9, 8), {'111': 200, '103a': 100, '109': 100, '108': 200,
                                     '107': 200, '110': 250, 'N_arumi01': 27}, 'TSSU5227711'),
        # CAAU8818353 9/3出荷・9/14着（PC多機能L 607）
        (datetime.date(2026, 9, 14), {'301': 208, '306': 399}, 'CAAU8818353'),
    ]
    # 工場の上线计划（品番別内訳が確定しているもの）
    with open(os.path.join(DATA, 'production-plans.json'), encoding='utf-8') as f:
        plans = json.load(f)['plans']
    for plan in plans:
        for sh in plan['shipments']:
            a.append((arrive(sh['ship']), dict(sh['items']), plan['order_no']))
    a.sort(key=lambda x: x[0])
    return a


def main():
    with open(os.path.join(DATA, 'inventory-suitcase-2026-09-04.csv'), encoding='utf-8-sig') as f:
        suit = list(csv.DictReader(f))
    stock = {r['SKU']: int(r['総在庫']) for r in suit}
    # 9/4一覧に無い品目は0（＝在庫登録待ちの可能性。注記で明示）
    for k in NOT_IN_SNAPSHOT | NEW_ITEMS:
        stock.setdefault(k, 0)
    # 入庫済みだがスナップショットに載っていない分を補正
    adj_path = os.path.join(DATA, 'stock-adjustments.json')
    adjusted = {}
    if os.path.exists(adj_path):
        with open(adj_path, encoding='utf-8') as f:
            for a in json.load(f)['adjustments']:
                stock[a['sku']] = stock.get(a['sku'], 0) + a['qty']
                adjusted[a['sku']] = a['qty']
    arrivals = build_arrivals(stock, DPS)

    dates = []
    d = STOCK_DATE
    while d <= CAL_END:
        dates.append(d); d += datetime.timedelta(days=1)
    arr_by_sku = {}
    for ad, items, _src in arrivals:
        for k, q in items.items():
            arr_by_sku.setdefault(k, {})
            arr_by_sku[k][ad] = arr_by_sku[k].get(ad, 0) + q

    # 日次シミュレーション
    cur_stock = dict(stock)
    hist = {k: {} for k in DPS}
    first_out, oos_days = {}, {k: 0 for k in DPS}
    for dt in dates:
        for k, q in arr_by_sku.items():
            if dt in q:
                cur_stock[k] = cur_stock.get(k, 0) + q[dt]
        for k in DPS:
            d0 = DPS[k]
            if d0 > 0:
                cur_stock[k] = max(0, cur_stock.get(k, 0) - d0)
                if cur_stock[k] <= 0:
                    oos_days[k] += 1
                    first_out.setdefault(k, dt)
            hist[k][dt] = round(cur_stock.get(k, 0))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '在庫消化カレンダー'
    ws['A1'] = f'📅 在庫消化カレンダー（{STOCK_DATE:%Y/%-m/%-d}時点・7月実績日販ベース）'
    ws['A1'].font = TF
    ws['A2'] = '凡例:'; ws['A2'].font = Font(bold=True)
    for i, (lab, fill) in enumerate([('赤=欠品', RED), ('緑=入荷日', GRN), ('青=在庫あり', BLU),
                                     ('黄=売切間近(7日分以下)', YEL), ('灰=販売なし', GRY)]):
        y = ws.cell(2, 2 + i * 2, lab)
        y.fill = fill; y.alignment = C; y.border = BD
        if fill in (RED, GRN):
            y.font = Font(bold=True, color='FFFFFF')
    ws['A3'] = ('日販=7月の3媒体（Amazon・楽天・自社）実績÷31日。'
                '現庫=9/4在庫（FBA合計＋CS倉庫良品＋FBM＋事務所）。'
                '入荷=確定便はパッキングリストの色別内訳、未出荷分は工場の上线计划（品番別）を出荷予定日+14日で計上。')
    ws['A3'].font = SUBF

    HEAD = ['SKU（色）', '日販', f'現庫({STOCK_DATE:%-m/%-d})', '入荷予定合計', '合計(現庫+入荷)',
            '欠品開始(予測)', '欠品日数', '在庫月数']
    BASE = len(HEAD) + 1  # 8列＋空列
    HROW = 6
    for c, h in enumerate(HEAD, 1):
        y = ws.cell(HROW, c, h); y.fill = HF; y.font = HFo; y.alignment = C; y.border = BD
    # 月ヘッダーは月ごとに結合して交互色で塗る（月の切れ目を見やすくする）
    month_start = {}
    for i, dt in enumerate(dates):
        month_start.setdefault((dt.year, dt.month), BASE + 1 + i)
    months = sorted(month_start)
    for mi, ym in enumerate(months):
        c0 = month_start[ym]
        c1 = month_start[months[mi + 1]] - 1 if mi + 1 < len(months) else BASE + len(dates)
        ws.merge_cells(start_row=HROW - 1, start_column=c0, end_row=HROW - 1, end_column=c1)
        fill = MON_A if mi % 2 == 0 else MON_B
        for c in range(c0, c1 + 1):
            y = ws.cell(HROW - 1, c); y.fill = fill; y.border = MONBD
        y = ws.cell(HROW - 1, c0, f'{ym[0]}年{ym[1]}月')
        y.font = Font(bold=True, size=11, color='FFFFFF'); y.alignment = C
    ws.row_dimensions[HROW - 1].height = 20
    for i, dt in enumerate(dates):
        col = BASE + 1 + i
        y = ws.cell(HROW, col, dt.day)
        y.fill = HF; y.font = Font(bold=True, color='FFFFFF', size=8); y.alignment = C
        y.border = MSTART if dt.day == 1 else BD
        ws.column_dimensions[get_column_letter(col)].width = 3.2

    r = HROW + 1
    for label, members in GROUPS:
        ws.cell(r, 1, label)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
        for c in range(1, len(HEAD) + 1):
            y = ws.cell(r, c); y.fill = GF; y.font = GFo; y.alignment = LN
        r += 1
        for disp, k in members:
            dps = DPS.get(k, 0.0)
            st = stock.get(k, 0)
            inc = sum(arr_by_sku.get(k, {}).values())
            fo = first_out.get(k)
            months = round((st + inc) / (dps * 30.4), 1) if dps > 0 else None
            note = '要確認' if k in NOT_IN_SNAPSHOT else ('新商品' if k in NEW_ITEMS else st)
            vals = [
                '　' + disp, dps, note, inc, (st + inc),
                fo.strftime('%-m/%-d') if fo else ('販売なし' if dps == 0 else f'{CAL_END:%Y/%-m}末まで無し'),
                oos_days[k] if dps > 0 else '—',
                f'{months}ヶ月' if months is not None else '—',
            ]
            for c, x in enumerate(vals, 1):
                y = ws.cell(r, c, x); y.border = BD; y.alignment = LN if c == 1 else C
                if c == 3 and k in NOT_IN_SNAPSHOT:
                    y.fill = YEL
                elif c == 3 and k in NEW_ITEMS:
                    y.fill = GRN
                if c == 6 and fo:
                    y.fill = RED; y.font = Font(bold=True, color='FFFFFF')
                elif c == 6 and dps == 0:
                    y.fill = GRY
                if c == 7 and dps > 0 and oos_days[k]:
                    y.fill = RED; y.font = Font(bold=True, color='FFFFFF')
            for i, dt in enumerate(dates):
                col = BASE + 1 + i
                val = hist[k].get(dt, 0)
                y = ws.cell(r, col, val); y.alignment = C; y.font = Font(size=8)
                if dt.day == 1:
                    y.border = MSTART
                if dps == 0:
                    y.fill = GRY
                elif dt in arr_by_sku.get(k, {}):
                    y.fill = GRN; y.font = Font(size=8, bold=True, color='FFFFFF')
                elif val <= 0:
                    y.fill = RED; y.font = Font(size=8, bold=True, color='FFFFFF')
                elif val < dps * 7:
                    y.fill = YEL
                else:
                    y.fill = BLU
            r += 1
    r += 1
    for note in [
        '※ 到着日=販売可能日として計算（大阪港到着→倉庫入庫のリードタイムぶん、実際は数日後ろにずれます）。',
        '※ 未出荷分の到着は「出荷予定日＋14日」で仮置き（実績平均）。着日が確定したら修正してください。',
        '※ クラシックアルミM（シルバー326／ブラック264＝590個・WHSU5830223 8/23入庫）は'
        '9/4在庫一覧に品目が無いため、data/stock-adjustments.json で現庫に加算しています。',
        '※ ジップS/M（計1,687個・EITU1030111＋CAAU5731521 8/25入庫）も同様に加算しています。'
        '次回の在庫スナップショットに載ったら data/stock-adjustments.json の該当分を削除してください（二重計上防止）。',
        '※ 反映済みの上线计划: LM20260618(3,500個)/LM20260625スポーツ(600個)/LM20260704 PC(5,600個)・アルミ(1,050個)。',
        '※ サザンモデル(LM20260907)は10/20・10/25・10/30出荷の3便=2,850個を計上。新商品のため現庫0・日販未設定。',
        '※ LM20260808は全13コンテナ(11/10〜2027/1/30出荷)を計上。色別合計 L2,600/M3,750/S4,400で検算一致。',
        '※ 日販は7月実績の平均。8月以降のイベントや繁忙期（年末年始・GW）の増減は織り込んでいません。',
    ]:
        ws.cell(r, 1, note).font = SUBF; r += 1

    ws.column_dimensions['A'].width = 32
    for col, w in (('B', 7), ('C', 11), ('D', 12), ('E', 14), ('F', 16), ('G', 9), ('H', 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(HROW + 1, BASE + 1)

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, '在庫消化カレンダー.xlsx')
    wb.save(path)
    print('saved', path)
    print(f'{"SKU":26s}{"日販":>6s}{"現庫":>7s}{"入荷":>7s}{"欠品開始":>10s}{"欠品日数":>8s}')
    for label, members in GROUPS:
        for disp, k in members:
            fo = first_out.get(k)
            print(f'{disp:26s}{DPS.get(k,0):6.2f}{stock.get(k,0):7d}'
                  f'{sum(arr_by_sku.get(k,{}).values()):7d}'
                  f'{(fo.strftime("%-m/%-d") if fo else "—"):>10s}{oos_days[k]:8d}')


if __name__ == '__main__':
    main()
