#!/usr/bin/env python3
"""生産計画の日中対照表（日本↔中国の認識合わせ用）を出力する。

  python3 scripts/buildPlanConfirm.py
  → out/生産計画_日中対照.txt   （LINE/メール貼り付け用）
  → out/生産計画_日中対照.xlsx  （日中を左右に並べた確認用シート）

書式（1行＝1色×1便）:
  注文番号 商品名:サイズ:カラー:数量:出荷日:余り

余りは色ごとの合計なので、その色の最終便の行にだけ数字を入れ、
それ以外の行は「－」とする（二重計上を避けるため）。
"""
import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'data')
OUT = os.path.join(ROOT, 'out')

# SKU → (日本語 商品名, サイズ, カラー) / (中文 品名, 尺寸, 颜色)
JP_SIZE = {'S': 'S', 'M': 'M', 'L': 'L', 'M+S': 'M+S'}
CN_SIZE = {'S': '20寸', 'M': '24寸', 'L': '28寸', 'M+S': '26寸'}

PC_COLORS = {
    '106': ('マットブラック', '磨砂黑色'), '107': ('マットシルバー', '磨砂银色'),
    '108': ('マットホワイト', '磨砂白色'), '110': ('マットグレー', '磨砂灰色'),
    '101': ('エナメルシルバー', '镜面银色'), '105': ('エナメルホワイト', '镜面白色'),
    '109': ('エナメルピンク', '镜面粉色'), '111': ('エナメルカーキ', '镜面卡其色'),
    '103a': ('エナメルスカイブルー', '镜面天蓝色'),
    '206': ('マットブラック', '磨砂黑色'), '207': ('マットシルバー', '磨砂银色'),
    '208': ('マットホワイト', '磨砂白色'), '210': ('マットグレー', '磨砂灰色'),
    '201': ('エナメルシルバー', '镜面银色'), '211': ('エナメルカーキ', '镜面卡其色'),
    '306': ('マットブラック', '磨砂黑色'), '307': ('マットシルバー', '磨砂银色'),
    '308': ('マットホワイト', '磨砂白色'), '310': ('マットグレー', '磨砂灰色'),
    '301': ('エナメルシルバー', '镜面银色'), 'L_khaki': ('エナメルカーキ', '镜面卡其色'),
}
SIZE_OF = {**{k: 'S' for k in ('106', '107', '108', '110', '101', '105', '109', '111', '103a')},
           **{k: 'M' for k in ('206', '207', '208', '210', '201', '211')},
           **{k: 'L' for k in ('306', '307', '308', '310', '301', 'L_khaki')}}

SPECIAL = {
    'arumi01': ('多機能アルミニウム', '多功能铝箱(前开盖铝箱)', 'S', 'シルバー', '银色'),
    'arumi02': ('多機能アルミニウム', '多功能铝箱(前开盖铝箱)', 'S', 'ブラック', '黑色'),
    'N_arumi01': ('クラシックアルミニウム', '拼接箱(全铝镁合金)', 'S', 'シルバー', '银色'),
    'N_arumi02': ('クラシックアルミニウム', '拼接箱(全铝镁合金)', 'S', 'ブラック', '黑色'),
    'outdoorsk001': ('スポーツモデル', '拼接箱26寸运动款', 'M+S', 'ブラック', '黑色'),
    'outdoorsk002': ('スポーツモデル', '拼接箱26寸运动款', 'M+S', 'シルバー', '银色'),
}
for sz, jsz in (('S', 'S'), ('M', 'M'), ('L', 'L')):
    for key, jp, cn in (('black', 'ブラック', '黑色'), ('gray', 'トグレー', '灰色'),
                        ('silver', 'シルバー', '银色'), ('white', 'ホワイト', '白色'),
                        ('turquoise', 'ターコイズ', '浅蓝色')):
        SPECIAL[f'TM_{sz}_{key}'] = ('サザンモデル', 'TM窄框款PC箱', jsz, jp, cn)

PC_NAME_JP, PC_NAME_CN = '多機能スーツケース', '多功能PC箱/前开PC箱'
CLASSIC_JP, CLASSIC_CN = '経典PC箱', '经典PC箱'


def describe(sku, order_no):
    if sku in SPECIAL:
        jp, cn, sz, cjp, ccn = SPECIAL[sku]
        return jp, cn, JP_SIZE[sz], CN_SIZE[sz], cjp, ccn
    cjp, ccn = PC_COLORS[sku]
    sz = SIZE_OF[sku]
    jp = CLASSIC_JP if order_no == 'LM20260808' else PC_NAME_JP
    cn = CLASSIC_CN if order_no == 'LM20260808' else PC_NAME_CN
    return jp, cn, JP_SIZE[sz], CN_SIZE[sz], cjp, ccn


def build_xlsx(plans, path):
    """オーダーごとに1シート。先頭に「サマリー」シートを置く。"""
    thin = Side(style='thin', color='B0B0B0')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    HF = PatternFill('solid', fgColor='1F4E78')
    HFo = Font(bold=True, color='FFFFFF', size=10)
    TF = Font(bold=True, size=14, color='1F4E78')
    SUB = Font(size=9, color='808080')
    SHIP = PatternFill('solid', fgColor='DDEBF7')
    TOT = PatternFill('solid', fgColor='F8CBAD')
    SPARE = PatternFill('solid', fgColor='FFF2CC')
    C = Alignment(horizontal='center', vertical='center')
    LN = Alignment(horizontal='left', vertical='center')
    WRAP = Alignment(horizontal='center', vertical='center', wrap_text=True)

    wb = openpyxl.Workbook()

    # ===== サマリー =====
    ws = wb.active
    ws.title = 'サマリー'
    ws['A1'] = 'オーダー別サマリー / 订单汇总'; ws['A1'].font = TF
    ws['A2'] = ('各オーダーの明細は下のシートタブに分けています。'
                'サイズ対応: S=20寸 / M=24寸 / L=28寸 / M+S=26寸')
    ws['A2'].font = SUB
    ws['A3'] = '各订单的明细请见下方各个工作表标签。'
    ws['A3'].font = SUB
    h = ['注文番号\n订单号', '内容', '数量\n数量', '出荷\n出货', '余り\n余数',
         '便数\n班次', '初回出荷\n首班', '最終出荷\n末班', '納期 / 交期']
    for c, x in enumerate(h, 1):
        y = ws.cell(5, c, x); y.fill = HF; y.font = HFo; y.border = bd; y.alignment = WRAP
    ws.row_dimensions[5].height = 30
    r = 6
    TQ = TS = TSP = TN = 0
    for pl in plans:
        tq = sum(pl['order_qty'].values())
        ts = sum(sh['qty'] for sh in pl['shipments'])
        tsp = sum(pl['spare_by_sku'].values())
        ships = [sh['ship'] for sh in pl['shipments']]
        vals = [pl['order_no'], pl['label'], tq, ts, tsp, len(ships),
                ships[0], ships[-1], pl['deadline']]
        for c, v in enumerate(vals, 1):
            y = ws.cell(r, c, v); y.border = bd
            y.alignment = LN if c in (2, 9) else C
            if c == 5 and v:
                y.fill = SPARE; y.font = Font(bold=True)
        TQ += tq; TS += ts; TSP += tsp; TN += len(ships); r += 1
    ws.cell(r, 2, '総合計 / 总合计')
    ws.cell(r, 3, TQ); ws.cell(r, 4, TS); ws.cell(r, 5, TSP); ws.cell(r, 6, TN)
    for c in range(1, 10):
        y = ws.cell(r, c); y.fill = TOT; y.font = Font(bold=True, size=11)
        y.border = bd; y.alignment = C if c != 2 else LN
    for i, w in enumerate([16, 42, 10, 10, 9, 8, 12, 12, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A6'

    # ===== オーダーごとのシート =====
    hdr = ['便No.\n班次', '商品名（日本語）', '品名（中文）', 'サイズ', '尺寸',
           'カラー（日本語）', '颜色（中文）', '数量', '出荷日\n出货日', '余り\n余数']
    for pl in plans:
        no = pl['order_no']
        oq, spare = pl['order_qty'], pl['spare_by_sku']
        last = {}
        for sh in pl['shipments']:
            for sku in sh['items']:
                last[sku] = sh['ship']
        tq = sum(oq.values())
        ts = sum(sh['qty'] for sh in pl['shipments'])
        tsp = sum(spare.values())

        w = wb.create_sheet(no[:31])
        w['A1'] = f"【{no}】{pl['label']}"; w['A1'].font = TF
        w['A2'] = f"数量 {tq:,}個 ／ 出荷 {ts:,}個 ／ 余り {tsp:,}個　　納期: {pl['deadline']}"
        w['A2'].font = Font(bold=True, size=11)
        w['A3'] = f"数量 {tq:,}个 ／ 出货 {ts:,}个 ／ 余数 {tsp:,}个　　交期: {pl['deadline']}"
        w['A3'].font = SUB
        w['A4'] = ('※余りは色ごとの合計です。その色の最終便の行にのみ記載しています（「－」は記載なし）。'
                   ' 余数为每个颜色的合计，只写在该颜色最后一班的行。')
        w['A4'].font = SUB
        if pl.get('note'):
            w['A5'] = '※ ' + pl['note']; w['A5'].font = Font(size=9, color='C00000')

        hr = 7
        for c, x in enumerate(hdr, 1):
            y = w.cell(hr, c, x); y.fill = HF; y.font = HFo; y.border = bd; y.alignment = WRAP
        w.row_dimensions[hr].height = 30
        r = hr + 1
        for sh in pl['shipments']:
            w.cell(r, 1, f"No.{sh['no']}　出荷日 出货日 {sh['ship']}　{sh['qty']:,}個 / 个")
            w.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
            for c in range(1, 11):
                y = w.cell(r, c); y.fill = SHIP; y.font = Font(bold=True); y.alignment = LN
            r += 1
            for sku, q in sh['items'].items():
                jp, cn, jsz, csz, cjp, ccn = describe(sku, no)
                sp = spare.get(sku, 0) if last[sku] == sh['ship'] else '－'
                vals = [f"No.{sh['no']}", jp, cn, jsz, csz, cjp, ccn, q, sh['ship'], sp]
                for c, v in enumerate(vals, 1):
                    y = w.cell(r, c, v); y.border = bd
                    y.alignment = LN if c in (2, 3) else C
                    if c == 10 and isinstance(v, int) and v > 0:
                        y.fill = SPARE; y.font = Font(bold=True)
                r += 1
        w.cell(r, 2, '合計 / 合计'); w.cell(r, 8, ts); w.cell(r, 10, tsp)
        for c in range(1, 11):
            y = w.cell(r, c); y.fill = TOT; y.font = Font(bold=True); y.border = bd
            y.alignment = C if c != 2 else LN
        for i, ww in enumerate([11, 24, 26, 8, 8, 20, 14, 9, 13, 9], 1):
            w.column_dimensions[get_column_letter(i)].width = ww
        w.freeze_panes = w.cell(hr + 1, 1)

    wb.save(path)


def main():
    with open(os.path.join(DATA, 'production-plans.json'), encoding='utf-8') as f:
        plans = json.load(f)['plans']
    lines = []
    lines.append('生産計画 日中対照表 / 生产计划 中日对照表')
    lines.append('書式 / 格式: 注文番号 商品名:サイズ:カラー:数量:出荷日:余り')
    lines.append('　　　　　　  订单号　 品名:尺寸:颜色:数量:出货日:余数')
    lines.append('※余りは色ごとの合計です。その色の最終便の行にのみ記載し、他は「－」としています。')
    lines.append('※余数为每个颜色的合计。只写在该颜色最后一班的行，其他行为「－」。')
    lines.append('')

    for pl in plans:
        no = pl['order_no']
        oq = pl['order_qty']
        spare = pl['spare_by_sku']
        # 色ごとの最終便を判定
        last = {}
        for sh in pl['shipments']:
            for sku in sh['items']:
                last[sku] = sh['ship']
        tot_q = sum(oq.values())
        tot_s = sum(sh['qty'] for sh in pl['shipments'])
        tot_sp = sum(spare.values())

        lines.append('=' * 62)
        lines.append(f"【{no}】{pl['label']}")
        lines.append(f"　数量合計 {tot_q:,}個 ／ 出荷合計 {tot_s:,}個 ／ 余り合計 {tot_sp:,}個")
        lines.append(f"　数量合计 {tot_q:,}个 ／ 出货合计 {tot_s:,}个 ／ 余数合计 {tot_sp:,}个")
        lines.append(f"　納期 / 交期: {pl['deadline']}")
        lines.append('')

        lines.append('■ 日本語')
        for sh in pl['shipments']:
            lines.append(f"  〔No.{sh['no']}　出荷日 {sh['ship']}　{sh['qty']:,}個〕")
            for sku, q in sh['items'].items():
                jp, _cn, jsz, _csz, cjp, _ccn = describe(sku, no)
                sp = f"{spare.get(sku, 0)}" if last[sku] == sh['ship'] else '－'
                lines.append(f"  {no} {jp}:{jsz}:{cjp}:{q}:{sh['ship']}:{sp}")
            lines.append('')

        lines.append('■ 中文')
        for sh in pl['shipments']:
            lines.append(f"  〔No.{sh['no']}　出货时间 {sh['ship']}　{sh['qty']:,}个〕")
            for sku, q in sh['items'].items():
                _jp, cn, _jsz, csz, _cjp, ccn = describe(sku, no)
                sp = f"{spare.get(sku, 0)}" if last[sku] == sh['ship'] else '－'
                lines.append(f"  {no} {cn}:{csz}:{ccn}:{q}:{sh['ship']}:{sp}")
            lines.append('')
        lines.append('')

    total_q = sum(sum(p['order_qty'].values()) for p in plans)
    total_s = sum(sh['qty'] for p in plans for sh in p['shipments'])
    total_sp = sum(sum(p['spare_by_sku'].values()) for p in plans)
    lines.append('=' * 62)
    lines.append(f"【総合計 / 总合计】数量 {total_q:,}個 ／ 出荷 {total_s:,}個 ／ 余り {total_sp:,}個")

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, '生産計画_日中対照.txt')
    with open(path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    print('saved', path)
    xlsx = os.path.join(OUT, '生産計画_日中対照.xlsx')
    build_xlsx(plans, xlsx)
    print('saved', xlsx)
    print(f'総合計 数量{total_q:,} / 出荷{total_s:,} / 余り{total_sp:,}')


if __name__ == '__main__':
    main()
