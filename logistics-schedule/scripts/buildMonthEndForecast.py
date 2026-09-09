#!/usr/bin/env python3
"""月末在庫予測（2026/9末〜2027/4末）をエクセルで出力する。

  python3 scripts/buildMonthEndForecast.py
  → out/月末在庫予測.xlsx

  シート1「区分別」  カテゴリ別の月末在庫＋月別の入荷/販売/増減
  シート2「SKU別」    SKU単位の月末在庫（日販・現庫つき）

在庫の起点は 9/4 スナップショット＋stock-adjustments.json の補正。
入荷は確定便（パッキングリスト）＋工場の上线计划（出荷予定日＋14日で到着）。
"""
import csv, datetime, json, os, sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import buildBurndownCalendar as B  # noqa: E402

ROOT = B.ROOT
DATA = B.DATA
OUT = B.OUT

MONTH_ENDS = [datetime.date(2026, 9, 30), datetime.date(2026, 10, 31),
              datetime.date(2026, 11, 30), datetime.date(2026, 12, 31),
              datetime.date(2027, 1, 31), datetime.date(2027, 2, 28),
              datetime.date(2027, 3, 31), datetime.date(2027, 4, 30)]

SHORT = {'ノーマルアルミ S（在庫表では「クラシックアルミ」）': 'クラシックアルミ S',
         'クラシックアルミ M（ノーマルアルミM）': 'クラシックアルミ M',
         'サザンモデル(TM窄框款PC箱)／新商品': 'サザンモデル'}

thin = Side(style='thin', color='B0B0B0')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill('solid', fgColor='1F4E78')
HFo = Font(bold=True, color='FFFFFF', size=10)
TF = Font(bold=True, size=14, color='1F4E78')
SUB = Font(size=9, color='808080')
GF = PatternFill('solid', fgColor='4472C4')
GFo = Font(bold=True, color='FFFFFF', size=11)
TOT = PatternFill('solid', fgColor='F8CBAD')
PEAK = PatternFill('solid', fgColor='FFC000')
IN = PatternFill('solid', fgColor='E2EFDA')
OUTF = PatternFill('solid', fgColor='FCE4D6')
GRY = PatternFill('solid', fgColor='D9D9D9')
RED = PatternFill('solid', fgColor='FF7C80')
YEL = PatternFill('solid', fgColor='FFFF00')
C = Alignment(horizontal='center', vertical='center')
LN = Alignment(horizontal='left', vertical='center')
WRAP = Alignment(horizontal='center', vertical='center', wrap_text=True)


def simulate():
    with open(os.path.join(DATA, 'inventory-suitcase-2026-09-04.csv'), encoding='utf-8-sig') as f:
        suit = list(csv.DictReader(f))
    stock = {r['SKU']: int(r['総在庫']) for r in suit}
    for k in B.NOT_IN_SNAPSHOT | B.NEW_ITEMS:
        stock.setdefault(k, 0)
    adj = {}
    with open(os.path.join(DATA, 'stock-adjustments.json'), encoding='utf-8') as f:
        for a in json.load(f)['adjustments']:
            stock[a['sku']] = stock.get(a['sku'], 0) + a['qty']
            adj[a['sku']] = adj.get(a['sku'], 0) + a['qty']

    arr = {}
    for ad, items, _src in B.build_arrivals(stock, B.DPS):
        for k, q in items.items():
            arr.setdefault(k, {})
            arr[k][ad] = arr[k].get(ad, 0) + q

    start = dict(stock)
    cur = dict(stock)
    snap, inflow, sold = {}, {}, {}
    d = B.STOCK_DATE
    while d <= MONTH_ENDS[-1]:
        mk = f'{d.year}/{d.month}'
        for k, q in arr.items():
            if d in q:
                cur[k] = cur.get(k, 0) + q[d]
                inflow[mk] = inflow.get(mk, 0) + q[d]
        for k, dps in B.DPS.items():
            if dps > 0:
                take = min(cur.get(k, 0), dps)
                cur[k] = cur.get(k, 0) - take
                sold[mk] = sold.get(mk, 0) + take
        if d in MONTH_ENDS:
            snap[d] = {k: round(v) for k, v in cur.items()}
        d += datetime.timedelta(days=1)
    return start, adj, snap, inflow, sold


def main():
    start, adj, snap, inflow, sold = simulate()
    grp = {}
    for label, members in B.GROUPS:
        for _disp, sku in members:
            grp[sku] = SHORT.get(label.replace('■ ', ''), label.replace('■ ', ''))
    groups = []
    for label, _m in B.GROUPS:
        g = SHORT.get(label.replace('■ ', ''), label.replace('■ ', ''))
        if g not in groups:
            groups.append(g)

    wb = openpyxl.Workbook()

    # ===== シート1 区分別 =====
    ws = wb.active
    ws.title = '区分別'
    ws['A1'] = '月末在庫予測（2026年9月末 〜 2027年4月末）'; ws['A1'].font = TF
    ws['A2'] = ('起点=2026/9/4在庫＋入庫補正（クラシックアルミM 590個・ジップS/M 1,687個）。'
                '入荷=確定便のパッキングリスト＋工場の上线计划（出荷予定日＋14日で到着と仮定）。'); ws['A2'].font = SUB
    ws['A3'] = ('日販は7月実績（3媒体÷31日）。日販未設定の区分は横ばい計算のため、'
                '灰色で示しています（実態より多く出ます）。'); ws['A3'].font = SUB

    hdr = ['区分', '現庫\n(9/4)'] + [f'{e.month}月末' for e in MONTH_ENDS]
    r = 5
    for c, h in enumerate(hdr, 1):
        y = ws.cell(r, c, h); y.fill = HF; y.font = HFo; y.border = BD; y.alignment = WRAP
    ws.row_dimensions[r].height = 26
    r += 1
    tot = [0] * len(MONTH_ENDS); tot0 = 0
    nodps = {g for g in groups
             if all(B.DPS.get(s, 0) == 0 for s, gg in grp.items() if gg == g)}
    for g in groups:
        base = sum(start.get(s, 0) for s, gg in grp.items() if gg == g)
        row = [sum(snap[e].get(s, 0) for s, gg in grp.items() if gg == g) for e in MONTH_ENDS]
        if base == 0 and not any(row):
            continue
        ws.cell(r, 1, g + ('（日販未設定）' if g in nodps else '')).alignment = LN
        ws.cell(r, 1).border = BD
        y = ws.cell(r, 2, base); y.border = BD; y.alignment = C
        for i, v in enumerate(row):
            y = ws.cell(r, 3 + i, v); y.border = BD; y.alignment = C
            if g in nodps:
                y.fill = GRY
        tot0 += base
        for i, v in enumerate(row):
            tot[i] += v
        r += 1
    ws.cell(r, 1, '合計').alignment = LN
    ws.cell(r, 2, tot0)
    peak = max(tot)
    for i, v in enumerate(tot):
        y = ws.cell(r, 3 + i, v)
        y.fill = PEAK if v == peak else TOT
    for c in range(1, 3 + len(MONTH_ENDS)):
        y = ws.cell(r, c); y.font = Font(bold=True, size=11); y.border = BD
        if not y.fill or y.fill.fgColor.rgb == '00000000':
            y.fill = TOT
        y.alignment = C if c != 1 else LN
    r += 2

    ws.cell(r, 1, '■ 月別の入荷・販売・増減').font = Font(bold=True, size=12, color='1F4E78')
    r += 1
    labels = ['月', '入荷', '販売', '増減']
    for i, lab in enumerate(labels):
        y = ws.cell(r + i, 1, lab); y.fill = HF; y.font = HFo; y.border = BD; y.alignment = C
    for j, e in enumerate(MONTH_ENDS):
        mk = f'{e.year}/{e.month}'
        i_, s_ = inflow.get(mk, 0), round(sold.get(mk, 0))
        for i, v in enumerate([f'{e.year}/{e.month}', i_, s_, i_ - s_]):
            y = ws.cell(r + i, 3 + j, v); y.border = BD; y.alignment = C
            if i == 1: y.fill = IN
            if i == 2: y.fill = OUTF
            if i == 3:
                y.font = Font(bold=True, color='C00000' if isinstance(v, int) and v < 0 else '006100')
    r += 5
    for note in [
        '※ 3月・4月は入荷が0本です。現在の最終発注は LM20260808（2027/1/30出荷＝2/13着）。',
        '※ 生産＋輸送に2〜3ヶ月かかるため、次の発注は12月〜1月に出さないと5〜7月に品薄になります。',
        '※ 日販未設定（灰色）の区分は減らない計算です。日販が分かれば精度が上がります。',
    ]:
        ws.cell(r, 1, note).font = SUB; r += 1

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 10
    for i in range(len(MONTH_ENDS)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 10
    ws.freeze_panes = 'C6'

    # ===== シート2 SKU別 =====
    ws2 = wb.create_sheet('SKU別')
    ws2['A1'] = '月末在庫予測（SKU別）'; ws2['A1'].font = TF
    ws2['A2'] = ('🟥=在庫0（欠品）／🟨=7日分未満。灰色=日販未設定（横ばい計算）。'
                 '「補正」列は9/4在庫一覧に品目が無く、入庫実績から加算した数量です。'); ws2['A2'].font = SUB
    h2 = ['区分', 'SKU（色）', '日販', '現庫\n(9/4)', '補正'] + [f'{e.month}月末' for e in MONTH_ENDS]
    r = 4
    for c, h in enumerate(h2, 1):
        y = ws2.cell(r, c, h); y.fill = HF; y.font = HFo; y.border = BD; y.alignment = WRAP
    ws2.row_dimensions[r].height = 26
    r += 1
    for label, members in B.GROUPS:
        g = SHORT.get(label.replace('■ ', ''), label.replace('■ ', ''))
        first = True
        for disp, sku in members:
            dps = B.DPS.get(sku, 0)
            base = start.get(sku, 0)
            row = [snap[e].get(sku, 0) for e in MONTH_ENDS]
            if base == 0 and not any(row) and dps == 0:
                continue
            ws2.cell(r, 1, g if first else '').alignment = LN
            ws2.cell(r, 1).border = BD
            first = False
            for c, v in enumerate([disp, dps if dps else '－', base,
                                   adj.get(sku, '') or ''], 2):
                y = ws2.cell(r, c, v); y.border = BD
                y.alignment = LN if c == 2 else C
                if c == 5 and v:
                    y.fill = IN; y.font = Font(bold=True)
            for i, v in enumerate(row):
                y = ws2.cell(r, 6 + i, v); y.border = BD; y.alignment = C
                if dps == 0:
                    y.fill = GRY
                elif v <= 0:
                    y.fill = RED; y.font = Font(bold=True, color='FFFFFF')
                elif v < dps * 7:
                    y.fill = YEL
            r += 1
    for i, w in enumerate([26, 30, 8, 10, 8], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for i in range(len(MONTH_ENDS)):
        ws2.column_dimensions[get_column_letter(6 + i)].width = 10
    ws2.freeze_panes = 'F5'

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, '月末在庫予測.xlsx')
    wb.save(path)
    print('saved', path)
    for i, e in enumerate(MONTH_ENDS):
        print(f'  {e.year}/{e.month}末  {tot[i]:,}個')


if __name__ == '__main__':
    main()
